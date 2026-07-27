"""
Consolidated report — reads whatever today's scanner sheets already wrote
into trading_toolkit.xlsx and emails one beautifully formatted HTML summary
covering all of them (NSE main scan, BSE gainers, BSE 1000 + Ultra Picks).

WHY GENERIC PARSING INSTEAD OF IMPORTING EACH SCANNER'S FUNCTIONS DIRECTLY:
this script doesn't need to know market_scanner.py's internal schema (column
names, function signatures) - it only relies on the SAME visual convention
every scanner here already uses: a section title row (green fill), a header
row (navy fill) right below it, then data rows until a blank row. That
convention is shared by market_scanner.py / bse_gainer_scanner.py /
bse1000_scanner.py's write_table / write_gainers_table / _write_section
functions - so this reads correctly regardless of exactly which columns
any individual scanner decides to add or rename later.

Run this AFTER all three scanners have already written into the same
workbook (see the GitHub Actions workflow - it runs them in sequence first).

Setup (once):
    pip install openpyxl

Usage:
    python send_scan_report.py trading_toolkit.xlsx

Environment variables required:
    GMAIL_SENDER        - the Gmail address sending the report
    GMAIL_APP_PASSWORD  - Gmail app password (not your normal password)
    REPORT_RECIPIENT    - optional, defaults to lightsoul.n@gmail.com
"""

import os
import sys
import smtplib
from datetime import date
from email.message import EmailMessage

from openpyxl import load_workbook

SEC_FILL_HEX = "548235"   # section title row background (green)
HDR_FILL_HEX = "1F3864"   # header row background (navy)
DEFAULT_RECIPIENT = "lightsoul.n@gmail.com"

# Which sheet-name prefixes to look for today, in the order they should
# appear in the email. Each scanner writes a sheet named "<prefix>_DDMonYY".
SHEET_PREFIXES = [
    ("BSE1000", "BSE 1000 Scan"),
    ("BSEGain", "BSE Gainers (5%-20%)"),
    ("Scan", "NSE Main Scan"),
]

VERDICT_COLORS = {
    "STRONG SETUP": "#C6E0B4",
    "WATCHLIST": "#FFE699",
    "AVOID": "#F8CBAD",
}


def _fill_hex(cell):
    try:
        rgb = cell.fill.fgColor.rgb
        if rgb and len(rgb) == 8:
            return rgb[2:].upper()  # strip alpha prefix
    except Exception:
        pass
    return None


def extract_sections(ws):
    """Generic reader: walks a worksheet and pulls out every
    title -> headers -> data-rows section it finds, based on the shared
    fill-color convention across all scanners."""
    sections = []
    row = 1
    max_row = ws.max_row
    max_col = ws.max_column

    while row <= max_row:
        cell = ws.cell(row=row, column=1)
        fill = _fill_hex(cell)

        if fill == SEC_FILL_HEX and cell.value:
            title = str(cell.value)
            row += 1
            header_row_idx = row
            headers = []
            c = 1
            while c <= max_col:
                hcell = ws.cell(row=header_row_idx, column=c)
                if hcell.value is None:
                    break
                headers.append(str(hcell.value))
                c += 1
            row += 1

            data_rows = []
            while row <= max_row:
                first = ws.cell(row=row, column=1)
                first_fill = _fill_hex(first)
                if first.value is None or first_fill == SEC_FILL_HEX:
                    break
                record = {}
                for i, h in enumerate(headers, start=1):
                    record[h] = ws.cell(row=row, column=i).value
                data_rows.append(record)
                row += 1

            sections.append({"title": title, "headers": headers, "rows": data_rows})
        else:
            row += 1

    return sections


def find_todays_sheet(wb, prefix, today):
    target = f"{prefix}_{today.strftime('%d%b%y')}"
    return wb[target] if target in wb.sheetnames else None


def _table_html(headers, rows, max_rows=15, highlight_verdict=True):
    if not rows:
        return "<p style='color:#888;font-style:italic;margin:4px 0 16px;'>No rows in this section today.</p>"

    shown = rows[:max_rows]
    thead = "".join(f"<th style='padding:6px 10px;text-align:center;'>{h}</th>" for h in headers)
    body_rows = []
    for i, r in enumerate(shown):
        bg = "#F7F7F7" if i % 2 else "#FFFFFF"
        verdict_val = r.get("Verdict")
        cells = []
        for h in headers:
            v = r.get(h)
            v = "" if v is None else v
            cell_bg = bg
            if highlight_verdict and h == "Verdict" and verdict_val in VERDICT_COLORS:
                cell_bg = VERDICT_COLORS[verdict_val]
            cells.append(f"<td style='padding:5px 10px;text-align:center;background:{cell_bg};'>{v}</td>")
        body_rows.append(f"<tr>{''.join(cells)}</tr>")

    more_note = ""
    if len(rows) > max_rows:
        more_note = (f"<p style='color:#888;font-size:12px;margin:4px 0 16px;'>"
                     f"+ {len(rows) - max_rows} more rows in the attached workbook.</p>")

    return (
        "<table style='border-collapse:collapse;width:100%;font-family:Arial,sans-serif;"
        "font-size:12px;margin-bottom:6px;'>"
        f"<thead><tr style='background:#1F3864;color:#FFFFFF;'>{thead}</tr></thead>"
        f"<tbody>{''.join(body_rows)}</tbody></table>{more_note}"
    )


def build_html_report(all_sections, today, missing_sheets):
    style_header = (
        "background:linear-gradient(135deg,#0B1F3A,#1F3864);color:#FFFFFF;"
        "padding:28px 32px;font-family:Arial,sans-serif;"
    )
    parts = [f"""
    <div style="{style_header}">
        <div style="font-size:22px;font-weight:bold;letter-spacing:0.5px;">
            Daily Market Scan &mdash; {today.strftime('%d %B %Y')}
        </div>
        <div style="font-size:13px;color:#D9E2F3;margin-top:4px;">
            Post-close consolidated report &middot; NSE + BSE
        </div>
    </div>
    """]

    if missing_sheets:
        missing_list = ", ".join(missing_sheets)
        parts.append(
            f"<div style='background:#FFF3CD;color:#7A5A00;padding:10px 20px;"
            f"font-family:Arial,sans-serif;font-size:13px;'>"
            f"&#9888; Could not find today's sheet for: {missing_list} &mdash; "
            f"that scanner may have failed or not run. Check the Actions log.</div>"
        )

    for scanner_label, sections in all_sections:
        parts.append(
            f"<div style='padding:20px 32px 4px;font-family:Arial,sans-serif;'>"
            f"<div style='font-size:16px;font-weight:bold;color:#1F3864;"
            f"border-bottom:2px solid #1F3864;padding-bottom:6px;margin-bottom:12px;'>"
            f"{scanner_label}</div></div>"
        )
        if not sections:
            parts.append("<div style='padding:0 32px 12px;font-family:Arial,sans-serif;"
                          "color:#888;font-style:italic;'>No sections found in this sheet.</div>")
            continue
        for sec in sections:
            parts.append(
                f"<div style='padding:0 32px;font-family:Arial,sans-serif;'>"
                f"<div style='font-size:13px;font-weight:bold;color:#548235;margin:10px 0 6px;'>"
                f"{sec['title']}</div>"
                f"{_table_html(sec['headers'], sec['rows'])}"
                f"</div>"
            )

    parts.append("""
    <div style="padding:20px 32px 28px;font-family:Arial,sans-serif;font-size:11px;color:#999;
                border-top:1px solid #eee;margin-top:12px;">
        Full detail for every row is in the attached workbook. Chart-pattern labels and
        "likely reason for gain" tags are heuristic screens, not guarantees &mdash;
        verify anything you intend to trade against the actual chart and announcement.
        This is a decision-support tool, not financial advice.
    </div>
    """)

    return "<html><body style='margin:0;padding:0;background:#F4F4F4;'>" + "".join(parts) + "</body></html>"


def send_email(html_body, xlsx_path, subject):
    sender = os.environ.get("GMAIL_SENDER")
    password = os.environ.get("GMAIL_APP_PASSWORD")
    recipient = os.environ.get("REPORT_RECIPIENT", DEFAULT_RECIPIENT)

    if not sender or not password:
        print("[!] GMAIL_SENDER / GMAIL_APP_PASSWORD not set — cannot send email.")
        print("    Report HTML was still generated — see daily_report_preview.html")
        with open("daily_report_preview.html", "w", encoding="utf-8") as f:
            f.write(html_body)
        return False

    msg = EmailMessage()
    msg["Subject"] = subject
    msg["From"] = sender
    msg["To"] = recipient
    msg.set_content("Your email client doesn't render HTML — see the attached workbook.")
    msg.add_alternative(html_body, subtype="html")

    with open(xlsx_path, "rb") as f:
        msg.add_attachment(f.read(), maintype="application",
                            subtype="vnd.openxmlformats-officedocument.spreadsheetml.sheet",
                            filename=os.path.basename(xlsx_path))

    with smtplib.SMTP_SSL("smtp.gmail.com", 465) as smtp:
        smtp.login(sender, password)
        smtp.send_message(msg)

    print(f"Email sent to {recipient}")
    return True


def main(xlsx_path):
    today = date.today()
    wb = load_workbook(xlsx_path, data_only=True)

    all_sections = []
    missing_sheets = []

    for prefix, label in SHEET_PREFIXES:
        ws = find_todays_sheet(wb, prefix, today)
        if ws is None:
            missing_sheets.append(label)
            all_sections.append((label, []))
            continue
        sections = extract_sections(ws)
        all_sections.append((label, sections))
        print(f"{label}: found sheet '{ws.title}', {len(sections)} section(s) parsed.")

    html = build_html_report(all_sections, today, missing_sheets)
    subject = f"Daily Market Scan — {today.strftime('%d %b %Y')}"
    send_email(html, xlsx_path, subject)


if __name__ == "__main__":
    if len(sys.argv) != 2:
        print("Usage: python send_scan_report.py trading_toolkit.xlsx")
        sys.exit(1)
    main(sys.argv[1])
