"""
BSE Top Gainer Scanner with reliable company-name resolution and optional
Next-Day Explosive Move integration.

This version fixes the case where a BSE numeric scrip code was written into both
"Symbol" and "Name" columns. It resolves names in this order:

1. Company/security/scrip-name column from the supplied CSV files.
2. Company name already returned by market_scanner.analyze_symbol().
3. Persistent local cache: bse_company_name_cache.json.
4. Yahoo Finance metadata for the exact BSE ticker (<scrip code>.BO).
5. A clear "Name unavailable (<code>)" value instead of silently repeating the ID.

Usage:
    python bse_gainer_scanner.py trading_toolkit.xlsx bse_smallcap250.csv [more_bse_universe.csv ...]
"""

from __future__ import annotations

import json
import os
import re
import sys
import time
import urllib.parse
import urllib.request
import xml.etree.ElementTree as ET
from datetime import date
from pathlib import Path
from typing import Any, Dict, Iterable, List, Optional, Sequence, Tuple

import pandas as pd
import yfinance as yf
from openpyxl import load_workbook
from openpyxl.styles import Alignment, Font
from openpyxl.utils import get_column_letter

from market_scanner import (
    ALT_FILL,
    BLACK,
    BORDER,
    FONT,
    HDR_FILL,
    HDR_FONT,
    SEC_FILL,
    SEC_FONT,
    _SHARED_SESSION,
    analyze_universe_parallel,
    get_nifty_returns,
    score_row,
    verdict,
)

MIN_GAIN_PCT = 5.0
MAX_GAIN_PCT = 20.0
DEEPDIVE_CAP_SEGMENT = "BSE Smallcap"
EXPLOSIVE_ANALYSIS_LIMIT = 60

BSE_NAME_CACHE_FILE = "bse_company_name_cache.json"
NAME_LOOKUP_RETRIES = 2
NAME_LOOKUP_DELAY_SECONDS = 0.25

GAIN_HEADERS = [
    "Symbol", "Name", "Cap Segment", "Close", "% Chg", "Volume",
    "Turnover (Rs Cr)", "Vol Ratio (vs 3d avg)", "Vol Ratio (vs 20d avg)",
    "RSI(14)", "ADX(14)", "Score (0-9)", "Verdict",
]

DEEPDIVE_HEADERS = [
    "Symbol", "Name", "Close", "% Chg", "Turnover (Rs Cr)",
    "Vol Ratio (vs 3d avg)", "Vol Ratio (vs 20d avg)",
    "Trend Strength", "20d Slope %/day", "Trend R^2",
    "Trend Quality (ADX x R^2)", "Entry", "Stop Loss",
    "Target 1", "Target 2", "Likely Reason for Gain", "Recent Headlines",
]

SYMBOL_COLUMN_ALIASES = (
    "scrip code", "scripcode", "security code", "securitycode", "bse code",
    "bsecode", "symbol", "ticker", "security id", "securityid", "scrip id",
    "scripid", "code",
)

NAME_COLUMN_ALIASES = (
    "company name", "companyname", "security name", "securityname",
    "scrip name", "scripname", "issuer name", "issuername",
    "name of company", "nameofcompany", "company", "constituents",
    "long name", "longname", "short name", "shortname", "name",
)

SECTOR_COLUMN_ALIASES = (
    "industry", "sector", "industry new name", "industrynewname",
    "industry name", "industryname", "sector name", "sectorname",
)

INVALID_NAME_VALUES = {
    "", "nan", "none", "null", "n/a", "na", "-", "--", "not available",
    "unknown",
}

REASON_KEYWORDS = {
    "Dividend / Bonus / Split": ["dividend", "bonus share", "stock split", "special dividend"],
    "Order / Contract Win": [
        "order win", "bags order", "contract win", "wins order", "secures order",
        "receives order", "awarded contract", "letter of intent", "lois worth",
    ],
    "Results / Earnings": [
        "q1 results", "q2 results", "q3 results", "q4 results", "net profit",
        "quarterly results", "profit jumps", "profit rises", "profit surges",
        "beats estimates", "ebitda",
    ],
    "Block / Bulk Deal": ["block deal", "bulk deal", "promoter buys", "stake buy"],
    "Rating Upgrade / Brokerage Call": [
        "upgrades", "rating upgrade", "target price raised", "buy rating",
        "brokerage", "initiates coverage",
    ],
    "Stake Sale / M&A": [
        "acquire", "acquisition", "stake sale", "merger", "amalgamation",
        "joint venture", " jv ",
    ],
    "Fundraise / QIP / Rights Issue": [
        "qip", "rights issue", "preferential issue", "raises funds", "fund raise",
    ],
    "Regulatory / Govt / Approval": [
        "government order", "regulatory approval", "receives approval", "clearance",
        "patent granted",
    ],
    "Buyback": ["buyback"],
}


def _normalise_header(value: Any) -> str:
    return re.sub(r"[^a-z0-9]+", "", str(value).strip().lower())


def _clean_text(value: Any) -> str:
    if value is None:
        return ""
    text = str(value).strip()
    return "" if text.lower() in INVALID_NAME_VALUES else text


def _clean_symbol(value: Any) -> str:
    text = _clean_text(value).upper()
    text = re.sub(r"\.(BO|NS)$", "", text, flags=re.IGNORECASE)
    if re.fullmatch(r"\d+\.0+", text):
        text = text.split(".", 1)[0]
    return text


def _is_numeric_only(value: str) -> bool:
    return bool(re.fullmatch(r"[\d\s.,+-]+", value.strip()))


def is_valid_company_name(value: Any, symbol: Any = "") -> bool:
    name = _clean_text(value)
    sym = _clean_symbol(symbol)
    if not name:
        return False
    if _is_numeric_only(name):
        return False
    if _clean_symbol(name) == sym:
        return False
    if name.lower().startswith("name unavailable"):
        return False
    return len(name) >= 2


def _find_column(df: pd.DataFrame, aliases: Sequence[str]) -> Optional[str]:
    """Find a column using exact normalised aliases before safe fuzzy matching."""
    normalised_columns = {_normalise_header(column): str(column) for column in df.columns}

    for alias in aliases:
        match = normalised_columns.get(_normalise_header(alias))
        if match is not None:
            return match

    # Conservative fuzzy match: only aliases of at least six characters.
    for alias in aliases:
        normalised_alias = _normalise_header(alias)
        if len(normalised_alias) < 6:
            continue
        for normalised_column, original in normalised_columns.items():
            if normalised_alias in normalised_column:
                return original
    return None


def _prefer_name(current: Any, candidate: Any, symbol: str) -> str:
    current_text = _clean_text(current)
    candidate_text = _clean_text(candidate)
    if is_valid_company_name(current_text, symbol):
        return current_text
    if is_valid_company_name(candidate_text, symbol):
        return candidate_text
    return ""


def load_bse_universe(universe_paths: Sequence[str], verbose: bool = True) -> pd.DataFrame:
    """Load BSE universe CSVs while preserving and validating actual company names.

    Unlike market_scanner.load_universe(), this loader understands common BSE
    headers such as Scrip Name, Issuer Name, Security Name and Scrip Code. All
    columns are read as strings so six-digit scrip codes remain intact.
    """
    records: Dict[str, Dict[str, str]] = {}

    for path_text in universe_paths:
        path = Path(path_text)
        try:
            df = pd.read_csv(path, dtype=str, keep_default_na=False)
        except UnicodeDecodeError:
            df = pd.read_csv(path, dtype=str, keep_default_na=False, encoding="latin-1")

        symbol_column = _find_column(df, SYMBOL_COLUMN_ALIASES)
        name_column = _find_column(df, NAME_COLUMN_ALIASES)
        sector_column = _find_column(df, SECTOR_COLUMN_ALIASES)

        if symbol_column is None:
            raise ValueError(
                f"{path}: no BSE symbol/scrip-code column found. Columns: {list(df.columns)}"
            )

        filename = path.name.lower()
        if "bse" in filename and "smallcap" in filename:
            cap_segment = "BSE Smallcap"
        elif "smallcap" in filename:
            cap_segment = "NSE Smallcap"
        else:
            cap_segment = "Large/Midcap"

        if verbose:
            print(
                f"[BSE names] {path.name}: symbol column={symbol_column!r}, "
                f"name column={name_column!r}, sector column={sector_column!r}"
            )

        for _, row in df.iterrows():
            symbol = _clean_symbol(row.get(symbol_column))
            if not symbol:
                continue

            candidate_name = _clean_text(row.get(name_column)) if name_column else ""
            candidate_sector = _clean_text(row.get(sector_column)) if sector_column else ""

            existing = records.setdefault(
                symbol,
                {"symbol": symbol, "name": "", "sector": "", "cap_segment": cap_segment},
            )
            existing["name"] = _prefer_name(existing.get("name"), candidate_name, symbol)
            if not existing.get("sector") and candidate_sector:
                existing["sector"] = candidate_sector
            if cap_segment == "BSE Smallcap":
                existing["cap_segment"] = "BSE Smallcap"

    if not records:
        raise ValueError("No valid BSE symbols were loaded from the supplied CSV files.")

    universe = pd.DataFrame(records.values())
    unresolved = int((~universe.apply(lambda r: is_valid_company_name(r["name"], r["symbol"]), axis=1)).sum())
    if verbose:
        print(
            f"[BSE names] Loaded {len(universe)} unique symbols; "
            f"{unresolved} require metadata/cache fallback."
        )
    return universe


def _load_name_cache(cache_path: str = BSE_NAME_CACHE_FILE) -> Dict[str, str]:
    try:
        with open(cache_path, "r", encoding="utf-8") as handle:
            raw = json.load(handle)
    except (FileNotFoundError, json.JSONDecodeError, OSError):
        return {}

    clean: Dict[str, str] = {}
    if isinstance(raw, dict):
        for symbol, name in raw.items():
            symbol_text = _clean_symbol(symbol)
            if symbol_text and is_valid_company_name(name, symbol_text):
                clean[symbol_text] = _clean_text(name)
    return clean


def _save_name_cache(cache: Dict[str, str], cache_path: str = BSE_NAME_CACHE_FILE) -> None:
    clean = {
        _clean_symbol(symbol): _clean_text(name)
        for symbol, name in cache.items()
        if _clean_symbol(symbol) and is_valid_company_name(name, symbol)
    }
    temporary_path = f"{cache_path}.tmp"
    with open(temporary_path, "w", encoding="utf-8") as handle:
        json.dump(dict(sorted(clean.items())), handle, ensure_ascii=False, indent=2)
    os.replace(temporary_path, cache_path)


def _name_from_mapping(mapping: Any, symbol: str) -> Optional[str]:
    if not isinstance(mapping, dict):
        return None
    for key in ("longName", "shortName", "displayName", "companyName", "name"):
        value = mapping.get(key)
        if is_valid_company_name(value, symbol):
            return _clean_text(value)
    return None


def fetch_bse_company_name(symbol: str, verbose: bool = True) -> Optional[str]:
    """Fetch the company name for the exact BSE ticker, never the NSE fallback."""
    clean_symbol = _clean_symbol(symbol)
    ticker_symbol = f"{clean_symbol}.BO"

    for attempt in range(1, NAME_LOOKUP_RETRIES + 1):
        try:
            ticker = yf.Ticker(ticker_symbol, session=_SHARED_SESSION)

            # Preferred metadata path.
            info = None
            get_info = getattr(ticker, "get_info", None)
            if callable(get_info):
                try:
                    info = get_info()
                except Exception:
                    info = None
            if not isinstance(info, dict):
                try:
                    info = ticker.info
                except Exception:
                    info = None

            name = _name_from_mapping(info, clean_symbol)
            if name:
                return name

            # Some yfinance versions populate history_metadata only after history().
            try:
                ticker.history(period="5d", interval="1d", auto_adjust=False)
            except Exception:
                pass
            name = _name_from_mapping(getattr(ticker, "history_metadata", None), clean_symbol)
            if name:
                return name

            # Last Yahoo fallback: exact-symbol search, available in newer yfinance.
            search_class = getattr(yf, "Search", None)
            if search_class is not None:
                try:
                    search = search_class(ticker_symbol, max_results=8)
                    quotes = getattr(search, "quotes", []) or []
                    for quote in quotes:
                        quote_symbol = str(quote.get("symbol", "")).upper()
                        if quote_symbol == ticker_symbol.upper():
                            name = _name_from_mapping(quote, clean_symbol)
                            if name:
                                return name
                except Exception:
                    pass
        except Exception as exc:
            if verbose:
                print(
                    f"[BSE names] {ticker_symbol}: lookup attempt {attempt}/"
                    f"{NAME_LOOKUP_RETRIES} failed ({exc})"
                )

        if attempt < NAME_LOOKUP_RETRIES:
            time.sleep(NAME_LOOKUP_DELAY_SECONDS * attempt)

    return None


def resolve_company_name(
    symbol: str,
    universe_name: Any,
    analysis_name: Any,
    cache: Dict[str, str],
    verbose: bool = True,
) -> str:
    """Resolve one BSE company name without ever silently returning the scrip ID."""
    clean_symbol = _clean_symbol(symbol)

    for candidate, source in (
        (universe_name, "CSV"),
        (analysis_name, "market scanner metadata"),
        (cache.get(clean_symbol), "local cache"),
    ):
        if is_valid_company_name(candidate, clean_symbol):
            name = _clean_text(candidate)
            cache[clean_symbol] = name
            if verbose and source != "CSV":
                print(f"[BSE names] {clean_symbol}: resolved from {source}: {name}")
            return name

    fetched_name = fetch_bse_company_name(clean_symbol, verbose=verbose)
    if is_valid_company_name(fetched_name, clean_symbol):
        name = _clean_text(fetched_name)
        cache[clean_symbol] = name
        if verbose:
            print(f"[BSE names] {clean_symbol}: resolved from {clean_symbol}.BO metadata: {name}")
        return name

    if verbose:
        print(
            f"[BSE names] WARNING: company name could not be resolved for {clean_symbol}. "
            f"Check that the input CSV contains Security Name/Scrip Name/Company Name."
        )
    return f"Name unavailable ({clean_symbol})"


def get_gain_reason(company_name: str, max_headlines: int = 6, timeout: int = 10) -> Tuple[str, List[str]]:
    query = urllib.parse.quote(f'"{company_name}" share')
    url = f"https://news.google.com/rss/search?q={query}%20when%3A3d&hl=en-IN&gl=IN&ceid=IN:en"
    try:
        request = urllib.request.Request(url, headers={"User-Agent": "Mozilla/5.0"})
        with urllib.request.urlopen(request, timeout=timeout) as response:
            data = response.read()
        root = ET.fromstring(data)
        titles = [
            item.find("title").text
            for item in root.iter("item")
            if item.find("title") is not None and item.find("title").text
        ][:max_headlines]
    except Exception:
        return "Could not fetch news (no internet or Google News blocked it)", []

    if not titles:
        return (
            "No recent headlines found — verify BSE announcements manually; "
            "may be a pure technical/momentum breakout",
            [],
        )

    combined = " | ".join(titles).lower()
    matched = [
        reason
        for reason, keywords in REASON_KEYWORDS.items()
        if any(keyword in combined for keyword in keywords)
    ]
    if not matched:
        return (
            "No obvious corporate-action keyword in headlines — likely technical/momentum move; "
            "verify manually",
            titles,
        )
    return " / ".join(matched), titles


def run_gainer_scan(universe_paths: Sequence[str], verbose: bool = True) -> Tuple[List[dict], List[dict]]:
    universe = load_bse_universe(universe_paths, verbose=verbose)
    symbols = universe["symbol"].tolist()
    meta = universe.set_index("symbol").to_dict(orient="index")
    name_cache = _load_name_cache()

    if verbose:
        print(f"BSE universe loaded: {len(symbols)} symbols")

    nifty_1m, nifty_3m = get_nifty_returns()
    fetched = analyze_universe_parallel(symbols, nifty_1m, nifty_3m, verbose=verbose)

    gainers: List[dict] = []
    for raw_symbol, result in fetched.items():
        symbol = _clean_symbol(raw_symbol)
        if result is None or result.get("pct_change") is None:
            continue
        if not (MIN_GAIN_PCT <= result["pct_change"] <= MAX_GAIN_PCT):
            continue

        row_meta = meta.get(symbol, {})
        result["symbol"] = symbol
        result["name"] = resolve_company_name(
            symbol=symbol,
            universe_name=row_meta.get("name"),
            analysis_name=result.get("company_name"),
            cache=name_cache,
            verbose=verbose,
        )
        result["sector"] = row_meta.get("sector", "")
        result["cap_segment"] = row_meta.get("cap_segment", "")
        result["score"] = score_row(result)
        result["verdict"] = verdict(result["score"])
        result["turnover_cr"] = (
            (result["close"] * result["volume"]) / 1e7
            if result.get("close") and result.get("volume")
            else None
        )
        gainers.append(result)

    try:
        _save_name_cache(name_cache)
    except OSError as exc:
        if verbose:
            print(f"[BSE names] Could not save {BSE_NAME_CACHE_FILE}: {exc}")

    gainers.sort(key=lambda item: item["pct_change"], reverse=True)

    deep_dive: List[dict] = []
    for result in gainers:
        if result.get("cap_segment") == DEEPDIVE_CAP_SEGMENT:
            reason, headlines = get_gain_reason(result["name"])
            result["gain_reason"] = reason
            result["headlines"] = "; ".join(headline[:90] for headline in headlines[:3]) if headlines else ""
            deep_dive.append(result)
            time.sleep(0.5)

    unresolved = [item["symbol"] for item in gainers if not is_valid_company_name(item["name"], item["symbol"])]
    if unresolved and verbose:
        print(f"[BSE names] Unresolved gainer names: {', '.join(unresolved)}")

    return gainers, deep_dive


def _num(value: Any, decimals: int = 2) -> Any:
    return round(value, decimals) if value is not None else "N/A"


def write_gainers_table(ws, start_row: int, rows: Sequence[dict], title: Optional[str] = None) -> int:
    row_number = start_row
    if title:
        ws.cell(row=row_number, column=1, value=title).font = SEC_FONT
        for column in range(1, len(GAIN_HEADERS) + 1):
            ws.cell(row=row_number, column=column).fill = SEC_FILL
        ws.merge_cells(start_row=row_number, start_column=1, end_row=row_number, end_column=len(GAIN_HEADERS))
        row_number += 1

    for column, header in enumerate(GAIN_HEADERS, start=1):
        cell = ws.cell(row=row_number, column=column, value=header)
        cell.font = HDR_FONT
        cell.fill = HDR_FILL
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        cell.border = BORDER
    ws.row_dimensions[row_number].height = 32
    row_number += 1

    for index, result in enumerate(rows):
        values = [
            str(result["symbol"]), str(result["name"]), result.get("cap_segment", ""),
            _num(result.get("close")), _num(result.get("pct_change")), result.get("volume"),
            _num(result.get("turnover_cr")), _num(result.get("vol_ratio_3d")),
            _num(result.get("vol_ratio_20d")), _num(result.get("rsi"), 1),
            _num(result.get("adx"), 1), result.get("score"), result.get("verdict"),
        ]
        for column, value in enumerate(values, start=1):
            cell = ws.cell(row=row_number, column=column, value=value)
            cell.font = BLACK
            cell.border = BORDER
            cell.alignment = Alignment(horizontal="left" if column == 2 else "center", wrap_text=(column == 2))
            if index % 2 == 1:
                cell.fill = ALT_FILL
        row_number += 1
    return row_number + 1


def write_gainer_deepdive_table(ws, start_row: int, rows: Sequence[dict], title: Optional[str] = None) -> int:
    row_number = start_row
    if title:
        ws.cell(row=row_number, column=1, value=title).font = SEC_FONT
        for column in range(1, len(DEEPDIVE_HEADERS) + 1):
            ws.cell(row=row_number, column=column).fill = SEC_FILL
        ws.merge_cells(start_row=row_number, start_column=1, end_row=row_number, end_column=len(DEEPDIVE_HEADERS))
        row_number += 1

    for column, header in enumerate(DEEPDIVE_HEADERS, start=1):
        cell = ws.cell(row=row_number, column=column, value=header)
        cell.font = HDR_FONT
        cell.fill = HDR_FILL
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        cell.border = BORDER
    ws.row_dimensions[row_number].height = 32
    row_number += 1

    for index, result in enumerate(rows):
        values = [
            str(result["symbol"]), str(result["name"]), _num(result.get("close")),
            _num(result.get("pct_change")), _num(result.get("turnover_cr")),
            _num(result.get("vol_ratio_3d")), _num(result.get("vol_ratio_20d")),
            result.get("trend_strength", "N/A"), _num(result.get("trend_slope_pct")),
            _num(result.get("trend_r2")), _num(result.get("trend_quality"), 1),
            _num(result.get("entry")), _num(result.get("stop_loss")),
            _num(result.get("target1")), _num(result.get("target2")),
            result.get("gain_reason", ""), result.get("headlines", ""),
        ]
        for column, value in enumerate(values, start=1):
            cell = ws.cell(row=row_number, column=column, value=value)
            cell.font = BLACK
            cell.border = BORDER
            cell.alignment = Alignment(
                horizontal="left" if column in (2, 16, 17) else "center",
                wrap_text=True,
            )
            if index % 2 == 1:
                cell.fill = ALT_FILL
        row_number += 1
    return row_number + 1


def _build_explosive_candidates(
    rows: Sequence[dict],
    nifty_1m: Optional[float],
    nifty_3m: Optional[float],
    market_regime: Any,
    exchange: str,
    ticker_suffix: str,
    scanner_name: str,
    limit: int = EXPLOSIVE_ANALYSIS_LIMIT,
):
    # Local imports prevent circular imports through explosive_engine -> market_scanner.
    from explosive_config import DEFAULT_CONFIG
    from explosive_engine import analyze_symbol_explosive

    explosive_candidates = []
    rejected = []

    for result in rows[:limit]:
        symbol = _clean_symbol(result.get("symbol"))
        if not symbol:
            continue
        ticker_symbol = symbol if symbol.endswith(ticker_suffix) else f"{symbol}{ticker_suffix}"
        try:
            history = yf.Ticker(ticker_symbol, session=_SHARED_SESSION).history(
                period=f"{DEFAULT_CONFIG.calibration_years}y",
                interval="1d",
                auto_adjust=True,
            )
        except Exception as exc:
            print(f"[Explosive] {symbol}: history fetch failed ({exc})")
            continue

        if history is None or history.empty:
            print(f"[Explosive] {symbol}: no multi-year history")
            continue
        if isinstance(history.columns, pd.MultiIndex):
            history.columns = history.columns.get_level_values(0)
        required_columns = ["Open", "High", "Low", "Close", "Volume"]
        if any(column not in history.columns for column in required_columns):
            print(f"[Explosive] {symbol}: missing required OHLCV columns")
            continue
        history = history.dropna(subset=["High", "Low", "Close", "Volume"])
        if history.empty:
            continue

        try:
            candidate = analyze_symbol_explosive(
                symbol=symbol,
                name=result.get("name", f"Name unavailable ({symbol})"),
                exchange=exchange,
                scanner_sources=[scanner_name],
                sector=result.get("sector", ""),
                df=history,
                nifty_1m=nifty_1m,
                nifty_3m=nifty_3m,
                market_regime=market_regime,
                config=DEFAULT_CONFIG,
                analysis_date=date.today(),
            )
        except Exception as exc:
            print(f"[Explosive] {symbol}: analysis failed ({exc})")
            continue

        (rejected if candidate.rejected else explosive_candidates).append(candidate)

    explosive_candidates.sort(key=lambda candidate: candidate.final_score, reverse=True)
    rejected.sort(key=lambda candidate: candidate.final_score, reverse=True)
    return explosive_candidates, rejected


def main(xlsx_path: str, universe_paths: Sequence[str]) -> None:
    # Local imports prevent unnecessary/circular loading during module import and tests.
    from explosive_market_regime import classify_regime
    from explosive_report import write_explosive_section

    gainers, deep_dive = run_gainer_scan(universe_paths)

    nifty_1m, nifty_3m = get_nifty_returns()
    benchmark_history = yf.Ticker("^BSESN", session=_SHARED_SESSION).history(
        period="1y", interval="1d", auto_adjust=True
    )
    if benchmark_history is None or benchmark_history.empty or "Close" not in benchmark_history.columns:
        raise RuntimeError("Could not fetch Sensex benchmark history for market-regime classification.")

    above_50dma_flags = [result.get("above_50dma", False) for result in gainers]
    regime = classify_regime(benchmark_history["Close"], above_50dma_flags)

    explosive_candidates, explosive_rejected = _build_explosive_candidates(
        rows=gainers,
        nifty_1m=nifty_1m,
        nifty_3m=nifty_3m,
        market_regime=regime,
        exchange="BSE",
        ticker_suffix=".BO",
        scanner_name="BSE Gainer Scanner",
    )

    today = date.today()
    sheet_name = f"BSEGain_{today.strftime('%d%b%y')}"

    workbook = load_workbook(xlsx_path)
    if sheet_name in workbook.sheetnames:
        del workbook[sheet_name]
    worksheet = workbook.create_sheet(sheet_name)
    worksheet.sheet_view.showGridLines = False
    worksheet.cell(
        row=1,
        column=1,
        value=f"BSE Top Gainers ({MIN_GAIN_PCT}%-{MAX_GAIN_PCT}%) — {today.strftime('%d %b %Y')}",
    ).font = Font(name=FONT, size=14, bold=True, color="1F3864")
    worksheet.merge_cells(start_row=1, start_column=1, end_row=1, end_column=len(GAIN_HEADERS))

    next_row = write_gainers_table(
        worksheet,
        3,
        gainers,
        title=f"ALL BSE GAINERS {MIN_GAIN_PCT}%-{MAX_GAIN_PCT}% — {len(gainers)} stocks",
    )
    next_row = write_gainer_deepdive_table(
        worksheet,
        next_row,
        deep_dive,
        title=f"BSE SMALLCAP GAINERS — DEEP DIVE + LIKELY REASON — {len(deep_dive)} stocks",
    )
    write_explosive_section(
        worksheet,
        next_row,
        explosive_candidates,
        explosive_rejected,
        [],
        {
            "Qualifying setups": len(explosive_candidates),
            "Rejected setups": len(explosive_rejected),
            "Market regime": regime.classification if hasattr(regime, "classification") else str(regime),
        },
        today.strftime("%d %b %Y"),
    )

    for column in range(1, max(len(GAIN_HEADERS), len(DEEPDIVE_HEADERS)) + 1):
        worksheet.column_dimensions[get_column_letter(column)].width = 15
    worksheet.column_dimensions["B"].width = 38
    worksheet.column_dimensions["P"].width = 44
    worksheet.column_dimensions["Q"].width = 54

    workbook.save(xlsx_path)
    print(
        f"\nDone. Sheet '{sheet_name}' written: {len(gainers)} total gainers, "
        f"{len(deep_dive)} BSE Smallcap deep-dive rows, "
        f"{len(explosive_candidates)} explosive candidates."
    )


if __name__ == "__main__":
    if len(sys.argv) < 3:
        print(
            "Usage: python bse_gainer_scanner.py trading_toolkit.xlsx "
            "bse_smallcap250.csv [more_bse_universe.csv ...]"
        )
        raise SystemExit(1)
    main(sys.argv[1], sys.argv[2:])
