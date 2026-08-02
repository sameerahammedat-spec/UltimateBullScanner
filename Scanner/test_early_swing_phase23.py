from __future__ import annotations

from dataclasses import replace
from datetime import date, datetime, timedelta, timezone

import pandas as pd
from openpyxl import Workbook, load_workbook

from early_swing_alerts import generate_alerts
from early_swing_backtest import _simulate_exit, score_band
from early_swing_calibration import (
    apply_calibration,
    build_calibration,
    estimate_candidate,
    load_calibration,
    save_calibration,
    wilson_interval,
)
from early_swing_config import DEFAULT_EARLY_SWING_CONFIG
from early_swing_engine import analyze_early_swing
from early_swing_features import classify_market_regime
from early_swing_models import BacktestTrade, UniverseRecord
from early_swing_report import write_early_swing_report
from test_early_swing import make_benchmark, make_breakout


def _record():
    return UniverseRecord("500001", "TEST.NS", "Test Ltd", "BSE", "A", "Industrials", ["BSE Group A"], "TEST")


def _candidate():
    df = make_breakout(age=1)
    benchmark = make_benchmark()
    candidate = analyze_early_swing(_record(), df, benchmark, analysis_date=df.index[-1].date())
    assert not candidate.rejected
    return candidate, benchmark


def _trade(index: int, result: str = "TARGET", setup_type: str = "Consolidation Breakout",
           score: float = 80.0, regime: str = "MILD_BULLISH") -> BacktestTrade:
    return BacktestTrade(
        symbol=f"S{index}", signal_date=f"2025-01-{(index % 28) + 1:02d}", entry_date="2025-02-01",
        entry_price=100.0, stop_price=95.0, target_price=106.0, exit_date="2025-02-05",
        exit_price=106.0 if result == "TARGET" else 95.0, result=result,
        return_pct=5.5 if result == "TARGET" else -5.5, holding_sessions=4,
        score=score, setup_type=setup_type, rally_age=1, market_regime=regime,
        score_band=score_band(score), entry_method="NEXT_OPEN", target_method="FIXED_PERCENT",
        holding_limit=10, gross_return_pct=6.0 if result == "TARGET" else -5.0,
        total_cost_pct=0.35, max_favourable_excursion_pct=7.0,
        max_adverse_excursion_pct=-2.0, atr_pct=3.0, liquidity_bucket="HIGH",
    )


def test_wilson_interval_is_bounded():
    low, high = wilson_interval(60, 100)
    assert 0 < low < 60 < high < 100


def test_calibration_uses_only_canonical_strategy():
    trades = [_trade(i, "TARGET" if i < 24 else "STOP") for i in range(30)]
    extra = _trade(99)
    extra.entry_method = "ABOVE_SIGNAL_HIGH"
    trades.append(extra)
    calibration = build_calibration(trades, replace(DEFAULT_EARLY_SWING_CONFIG, calibration_minimum_sample=10))
    assert calibration["selected_trade_count"] == 30
    global_bucket = next(row for row in calibration["buckets"] if row["key"] == "GLOBAL")
    assert global_bucket["sample_size"] == 30
    assert round(global_bucket["hit_rate_pct"], 1) == 80.0


def test_candidate_receives_empirical_hit_rate_with_fallback():
    candidate, _ = _candidate()
    config = replace(DEFAULT_EARLY_SWING_CONFIG, calibration_minimum_sample=10)
    trades = [_trade(i, "TARGET" if i < 21 else "STOP", setup_type=candidate.setup_type,
                     score=candidate.rules_score, regime=candidate.market_regime) for i in range(30)]
    calibration = build_calibration(trades, config)
    estimate = estimate_candidate(candidate, calibration, config)
    assert estimate.probability_label == "EMPIRICAL_HIT_RATE"
    assert estimate.sample_size >= 30
    assert estimate.hit_rate_pct == 70.0
    apply_calibration([candidate], calibration, config)
    assert candidate.empirical_hit_rate_pct == 70.0


def test_calibration_save_and_load(tmp_path):
    calibration = build_calibration([_trade(i) for i in range(30)], replace(DEFAULT_EARLY_SWING_CONFIG, calibration_minimum_sample=10))
    path = tmp_path / "calibration.json"
    save_calibration(calibration, path)
    loaded = load_calibration(path)
    assert loaded is not None
    assert loaded["selected_trade_count"] == 30


def test_alert_deduplicates_same_setup(tmp_path):
    candidate, _ = _candidate()
    candidate.rules_score = 90
    candidate.suggested_status = "Enter"
    candidate.current_price = candidate.entry_low
    config = replace(DEFAULT_EARLY_SWING_CONFIG, alert_minimum_rules_score=75, alert_cooldown_days=5)
    state = tmp_path / "state.json"
    output = tmp_path / "alerts.csv"
    now = datetime(2026, 8, 2, tzinfo=timezone.utc)
    first = generate_alerts([candidate], config, str(state), str(output), now=now)
    second = generate_alerts([candidate], config, str(state), str(output), now=now + timedelta(days=1))
    assert len(first) == 1
    assert first[0].alert_status == "READY_NEAR_ENTRY"
    assert second == []


def test_alert_emits_state_transition_to_do_not_chase(tmp_path):
    candidate, _ = _candidate()
    candidate.rules_score = 90
    candidate.suggested_status = "Enter"
    candidate.current_price = candidate.entry_low
    config = replace(DEFAULT_EARLY_SWING_CONFIG, alert_minimum_rules_score=75)
    state = tmp_path / "state.json"
    output = tmp_path / "alerts.csv"
    now = datetime(2026, 8, 2, tzinfo=timezone.utc)
    generate_alerts([candidate], config, str(state), str(output), now=now)
    candidate.current_price = candidate.entry_high * 1.10
    changed = generate_alerts([candidate], config, str(state), str(output), now=now + timedelta(hours=2))
    assert len(changed) == 1
    assert changed[0].alert_status == "ABOVE_ENTRY_DO_NOT_CHASE"


def test_conservative_exit_assumes_stop_first():
    frame = pd.DataFrame({
        "Open": [100.0], "High": [108.0], "Low": [94.0], "Close": [102.0], "Volume": [100000],
    }, index=pd.bdate_range("2026-01-01", periods=1))
    result, _, price, _, _ = _simulate_exit(frame, 0, 100.0, 95.0, 106.0, 1)
    assert result == "STOP"
    assert price == 95.0


def test_report_includes_phase2_and_phase3_sections(tmp_path):
    candidate, benchmark = _candidate()
    candidate.rules_score = 90
    candidate.suggested_status = "Enter"
    candidate.current_price = candidate.entry_low
    config = replace(DEFAULT_EARLY_SWING_CONFIG, calibration_minimum_sample=10)
    trades = [_trade(i, "TARGET" if i < 20 else "STOP", setup_type=candidate.setup_type,
                     score=candidate.rules_score, regime=candidate.market_regime) for i in range(30)]
    calibration = build_calibration(trades, config)
    apply_calibration([candidate], calibration, config)
    alerts = generate_alerts([candidate], config, str(tmp_path / "state.json"), str(tmp_path / "alerts.csv"),
                             now=datetime(2026, 8, 2, tzinfo=timezone.utc))
    xlsx = tmp_path / "report.xlsx"
    Workbook().save(xlsx)
    sheet = write_early_swing_report(str(xlsx), [candidate], classify_market_regime(benchmark), date(2026, 8, 2),
                                     alerts=alerts, calibration=calibration)
    wb = load_workbook(xlsx, read_only=True)
    ws = wb[sheet]
    values = [cell.value for row in ws.iter_rows() for cell in row if cell.value]
    assert "EMPIRICAL CALIBRATION SUMMARY" in values
    assert any(str(value).startswith("NEW / CHANGED ALERTS") for value in values)
    assert "Empirical Hit Rate %" in values


def test_backtest_prefilter_finds_generated_breakout():
    from early_swing_backtest import _candidate_signal_indices
    frame = make_breakout(age=1)
    indices = _candidate_signal_indices(frame, DEFAULT_EARLY_SWING_CONFIG)
    assert len(frame) - 1 in indices


def test_cache_rejects_short_history_for_long_calibration(tmp_path):
    import pickle
    from early_swing_scanner import _cache_path, _load_cached
    record = _record()
    cache_dir = tmp_path / "cache"
    cache_dir.mkdir()
    path = _cache_path(cache_dir, record.stable_key)
    with path.open("wb") as handle:
        pickle.dump({"source_ticker": "TEST.NS", "history_period": "1y", "data": make_breakout()}, handle)
    assert _load_cached(cache_dir, record, 8, "5y") is None
    assert _load_cached(cache_dir, record, 8, "1y") is not None


def test_alert_marks_disappeared_setup_invalidated(tmp_path):
    from copy import deepcopy
    candidate, _ = _candidate()
    candidate.rules_score = 90
    candidate.suggested_status = "Enter"
    candidate.current_price = candidate.entry_low
    config = replace(DEFAULT_EARLY_SWING_CONFIG, alert_minimum_rules_score=75, alert_setup_expiry_days=30)
    state = tmp_path / "state.json"
    output = tmp_path / "alerts.csv"
    now = datetime(2026, 8, 2, tzinfo=timezone.utc)
    generate_alerts([candidate], config, str(state), str(output), now=now)

    rejected = deepcopy(candidate)
    rejected.rejected = True
    rejected.rejection_reason = "No valid setup today"
    rejected.setup_date = ""
    rejected.setup_type = "N/A"
    rejected.pivot = None
    rejected.entry_low = None
    rejected.entry_high = None
    rejected.current_price = float(candidate.stop_loss) - 1.0
    changed = generate_alerts([rejected], config, str(state), str(output), now=now + timedelta(days=1))
    assert any(event.alert_status == "INVALIDATED" for event in changed)
