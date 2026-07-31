"""
Excel writer for the Next-Day Explosive Move Engine.

Reuses market_scanner.py's existing style constants (FONT, HDR_FILL,
SEC_FILL, ALT_FILL, BORDER) so the Explosive tables visually match every
other sheet in the workbook - Section 20 explicitly requires preserving
the existing professional style, not introducing a different look.
"""

from datetime import date
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.formatting.rule import CellIsRule
from openpyxl.utils import get_column_letter

from market_scanner import FONT, BLACK, HDR_FONT, HDR_FILL, SEC_FILL, SEC_FONT, ALT_FILL, BORDER

BAND_COLORS = {
    "A+ EXCEPTIONAL SETUP": PatternFill("solid", fgColor="1E7A34"),   # dark green
    "A HIGH-CONVICTION WATCH": PatternFill("solid", fgColor="C6E0B4"),  # light green
    "B+ STRONG WATCHLIST": PatternFill("solid", fgColor="FFE699"),    # yellow
    "B CONDITIONAL SETUP": PatternFill("solid", fgColor="F4B183"),    # orange
    "DO NOT INCLUDE IN PRIMARY LIST": PatternFill("solid", fgColor="F8CBAD"),  # red
}
GREY_FILL = PatternFill("solid", fgColor="D9D9D9")   # insufficient data

TABLE1_HEADERS = [
    "Rank", "Symbol", "Company", "Exchange", "Scanner Source", "Sector", "Close",
    "Primary Pattern", "Secondary Pattern", "Setup Stage", "Setup Duration",
    "Pivot", "Distance to Pivot %", "Breakout Trigger", "Setup Low", "Invalidation",
    "Risk to Invalidation %", "ATR %", "Five-Day Range Compression", "Bollinger Width Percentile",
    "Volume Dry-Up Ratio", "Current Relative Volume", "Close-Location Value",
    "Relative Strength 1M", "Relative Strength 3M", "Market Regime", "Circuit Risk",
    "Pattern Score", "Compression Score", "Trend/RS Score", "Volume Score", "Market Score",
    "Penalty", "Explosive Setup Score", "Historical Probability/Hit Rate", "Historical Sample Size",
    "Confidence", "Data Quality", "Positive Reasons", "Main Risks", "Invalidation Reason", "Final Verdict",
]

TABLE2_HEADERS = ["Symbol", "Company", "Close", "Setup Score (if any)", "Rejection Reason"]

TABLE3_HEADERS = ["Score Bucket", "Pattern", "Sample Count", "Hit Rate %",
                   "Avg Next-Day High Return %", "Median Next-Day Close Return %",
                   "Avg Adverse Excursion %"]

TABLE4_HEADERS = ["Metric", "Value"]


def _num(v, nd=2):
    return round(v, nd) if isinstance(v, (int, float)) else ("N/A" if v is None else v)


def _write_headers(ws, row, headers):
    for i, h in enumerate(headers, start=1):
        c = ws.cell(row=row, column=i, value=h)
        c.font = HDR_FONT
        c.fill = HDR_FILL
        c.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        c.border = BORDER
    ws.row_dimensions[row].height = 34


def write_table1_candidates(ws, start_row, candidates, title=None):
    r = start_row
    if title:
        ws.cell(row=r, column=1, value=title).font = SEC_FONT
        for c in range(1, len(TABLE1_HEADERS) + 1):
            ws.cell(row=r, column=c).fill = SEC_FILL
        ws.merge_cells(start_row=r, start_column=1, end_row=r, end_column=len(TABLE1_HEADERS))
        r += 1
    _write_headers(ws, r, TABLE1_HEADERS)
    header_row = r
    r += 1

    for idx, cand in enumerate(candidates):
        prob_display = (f"{cand.hit_rate_pct}%" if cand.probability_label == "EMPIRICAL_HIT_RATE"
                         and cand.hit_rate_pct is not None else "NOT_AVAILABLE")
        vals = [
            idx + 1, cand.symbol, cand.name, cand.exchange, ", ".join(cand.scanner_sources), cand.sector,
            _num(cand.close), cand.primary_pattern, cand.secondary_pattern or "-", cand.setup_stage,
            cand.setup_duration or "N/A", _num(cand.pivot), _num(cand.distance_to_pivot_pct),
            _num(cand.breakout_trigger), _num(cand.setup_low), _num(cand.invalidation),
            _num(cand.risk_to_invalidation_pct), _num(cand.atr_pct), _num(cand.range_compression_5_20),
            _num(cand.bb_width_percentile, 0), _num(cand.volume_dry_up_ratio), _num(cand.relative_volume),
            _num(cand.close_location_value), _num(cand.rs_1m), _num(cand.rs_3m), cand.market_regime,
            cand.circuit_risk, _num(cand.pattern_score), _num(cand.compression_score),
            _num(cand.trend_rs_score), _num(cand.volume_score), _num(cand.market_score),
            _num(cand.penalty_total), _num(cand.final_score), prob_display, cand.sample_size,
            cand.confidence, cand.data_quality_status, "; ".join(cand.positive_reasons),
            "; ".join(cand.main_risks), cand.invalidation_reason, cand.final_verdict,
        ]
        band_fill = BAND_COLORS.get(cand.final_verdict, ALT_FILL if idx % 2 else None)
        for c, v in enumerate(vals, start=1):
            cell = ws.cell(row=r, column=c, value=v)
            cell.font = BLACK
            cell.border = BORDER
            cell.alignment = Alignment(horizontal="center", wrap_text=True)
            if band_fill:
                cell.fill = band_fill
        r += 1
    return r + 1, header_row


def write_table2_rejected(ws, start_row, rejected_candidates, title=None):
    r = start_row
    if title:
        ws.cell(row=r, column=1, value=title).font = SEC_FONT
        for c in range(1, len(TABLE2_HEADERS) + 1):
            ws.cell(row=r, column=c).fill = SEC_FILL
        ws.merge_cells(start_row=r, start_column=1, end_row=r, end_column=len(TABLE2_HEADERS))
        r += 1
    _write_headers(ws, r, TABLE2_HEADERS)
    r += 1
    for idx, cand in enumerate(rejected_candidates):
        vals = [cand.symbol, cand.name, _num(cand.close), _num(cand.final_score) if cand.final_score else "N/A",
                cand.rejection_reason]
        for c, v in enumerate(vals, start=1):
            cell = ws.cell(row=r, column=c, value=v)
            cell.font = BLACK
            cell.border = BORDER
            cell.alignment = Alignment(horizontal="center", wrap_text=True)
            if idx % 2:
                cell.fill = ALT_FILL
            if cand.data_quality_status == "INSUFFICIENT_DATA":
                cell.fill = GREY_FILL
        r += 1
    return r + 1


def write_table3_backtest(ws, start_row, backtest_rows, title=None):
    r = start_row
    if title:
        ws.cell(row=r, column=1, value=title).font = SEC_FONT
        for c in range(1, len(TABLE3_HEADERS) + 1):
            ws.cell(row=r, column=c).fill = SEC_FILL
        ws.merge_cells(start_row=r, start_column=1, end_row=r, end_column=len(TABLE3_HEADERS))
        r += 1
    _write_headers(ws, r, TABLE3_HEADERS)
    r += 1
    for idx, row_data in enumerate(backtest_rows):
        for c, v in enumerate(row_data, start=1):
            cell = ws.cell(row=r, column=c, value=v)
            cell.font = BLACK
            cell.border = BORDER
            cell.alignment = Alignment(horizontal="center")
            if idx % 2:
                cell.fill = ALT_FILL
        r += 1
    return r + 1


def write_table4_dashboard(ws, start_row, dashboard_dict, title=None):
    r = start_row
    if title:
        ws.cell(row=r, column=1, value=title).font = SEC_FONT
        for c in range(1, 3):
            ws.cell(row=r, column=c).fill = SEC_FILL
        ws.merge_cells(start_row=r, start_column=1, end_row=r, end_column=2)
        r += 1
    _write_headers(ws, r, TABLE4_HEADERS)
    r += 1
    for idx, (k, v) in enumerate(dashboard_dict.items()):
        ws.cell(row=r, column=1, value=k).font = BLACK
        ws.cell(row=r, column=2, value=v).font = BLACK
        for c in (1, 2):
            ws.cell(row=r, column=c).border = BORDER
            if idx % 2:
                ws.cell(row=r, column=c).fill = ALT_FILL
        r += 1
    return r + 1


def write_explosive_section(ws, start_row, candidates, rejected, backtest_rows, dashboard, date_label):
    """Writes all 4 tables in sequence starting at start_row. Returns the
    next free row (for callers appending this after their own existing
    tables in the SAME sheet, per Section 20's 'add a new dated section'
    instruction - not necessarily a whole separate sheet)."""
    r = start_row
    ws.cell(row=r, column=1, value=f"NEXT-DAY EXPLOSIVE MOVE WATCHLIST — {date_label}").font = \
        Font(name=FONT, size=13, bold=True, color="1F3864")
    ws.merge_cells(start_row=r, start_column=1, end_row=r, end_column=len(TABLE1_HEADERS))
    r += 2

    r, _ = write_table1_candidates(ws, r, candidates,
        title=f"TOP NEXT-DAY EXPLOSIVE CANDIDATES — {len(candidates)} qualifying setups"
              if candidates else "TOP NEXT-DAY EXPLOSIVE CANDIDATES — NO QUALIFYING SETUPS TODAY")
    r = write_table2_rejected(ws, r, rejected,
        title=f"REJECTED HIGH-SCORE STOCKS — {len(rejected)} excluded")
    r = write_table3_backtest(ws, r, backtest_rows,
        title="BACKTEST SUMMARY (per-symbol empirical calibration, pooled)")
    r = write_table4_dashboard(ws, r, dashboard, title="MARKET DASHBOARD")

    for i in range(1, len(TABLE1_HEADERS) + 1):
        ws.column_dimensions[get_column_letter(i)].width = 14
    ws.column_dimensions["C"].width = 24
    return r
