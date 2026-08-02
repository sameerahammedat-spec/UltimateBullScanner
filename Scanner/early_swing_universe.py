"""Robust universe loading and deduplication for the Early Swing scanner.

BSE Group A/B exports sometimes contain trailing commas beyond the header. They
also contain three easily confused identifiers:

* Security Code: six-digit BSE code, e.g. 500002
* Security Id: trading symbol, e.g. ABB
* ISIN: e.g. INE117A01022 (never a Yahoo ticker)

This module keeps the Security Code as the stable identity, but makes Security Id
the preferred Yahoo ticker. Records are enriched during deduplication, so a BSE
1000 row containing only a numeric code can inherit the Security Id from Group A
or Group B before any network request is made.
"""
from __future__ import annotations

import csv
import re
from pathlib import Path
from typing import Dict, List, Optional, Sequence, Tuple

from early_swing_models import UniverseRecord


# Keep Security Code aliases separate from Security Id aliases. The former is the
# stable deduplication key; the latter is the preferred Yahoo symbol base.
SYMBOL_ALIASES = (
    "security code", "scrip code", "bse code", "securitycode", "scripcode",
    "bsecode", "symbol", "ticker",
)
NAME_ALIASES = (
    "security name", "issuer name", "company name", "constituents", "scrip name", "name"
)
SECTOR_ALIASES = (
    "macro-economic sector", "macro economic sector", "industry", "sector"
)
GROUP_ALIASES = ("group", "bse group")
SECURITY_ID_ALIASES = ("security id", "securityid", "scrip id", "scripid")
ISIN_ALIASES = ("isin", "isin no", "isin number")


def _norm(value: str) -> str:
    return re.sub(r"[^a-z0-9]+", "", str(value).strip().lower())


def _clean(value: object) -> str:
    text = "" if value is None else str(value).strip()
    return "" if text.lower() in {"", "nan", "none", "null", "n/a", "na", "-", "--"} else text


def _find_header(headers: Sequence[str], aliases: Sequence[str]) -> Optional[int]:
    normalised = [_norm(h) for h in headers]
    for alias in aliases:
        target = _norm(alias)
        if target in normalised:
            return normalised.index(target)
    for alias in aliases:
        target = _norm(alias)
        if len(target) < 5:
            continue
        for idx, header in enumerate(normalised):
            if target in header:
                return idx
    return None


def _source_label(path: Path) -> str:
    name = path.name.lower()
    if "group_a" in name or "groupa" in name or "group-a" in name:
        return "BSE Group A"
    if "group_b" in name or "groupb" in name or "group-b" in name:
        return "BSE Group B"
    if "1000" in name:
        return "BSE 1000"
    if "gainer" in name:
        return "BSE Gainer"
    if "full" in name and "bse" in name:
        return "BSE Full Universe"
    return path.stem


def _parse_label_path(value: str) -> Tuple[str, Path]:
    # Optional CLI form: "BSE Gainer=C:\\path\\gainers.csv".
    if "=" in value:
        label, candidate = value.split("=", 1)
        p = Path(candidate)
        if p.exists():
            return label.strip(), p
    p = Path(value)
    return _source_label(p), p


def _normalise_security_code(value: object) -> str:
    text = _clean(value)
    if re.fullmatch(r"\d+\.0+", text):
        text = text.split(".", 1)[0]
    return re.sub(r"\D", "", text) if text else ""


def _valid_security_id(value: object, symbol: str = "", isin: str = "") -> str:
    text = _clean(value).upper()
    text = re.sub(r"\.(BO|NS)$", "", text, flags=re.IGNORECASE)
    text = re.sub(r"[^A-Z0-9&-]+", "", text)
    if not text or text.isdigit():
        return ""
    if symbol and text == symbol:
        return ""
    if re.fullmatch(r"INE[A-Z0-9]{9}", text):
        return ""
    if isin and _norm(text) == _norm(isin):
        return ""
    return text


def _preferred_ticker(symbol: str, security_id: str, exchange: str) -> str:
    if exchange == "BSE":
        return f"{security_id}.BO" if security_id else f"{symbol}.BO"
    base = re.sub(r"\.(BO|NS)$", "", symbol.upper(), flags=re.IGNORECASE)
    return f"{base}.NS"


def load_universe_file(path_value: str) -> List[UniverseRecord]:
    label, path = _parse_label_path(path_value)
    if not path.exists():
        raise FileNotFoundError(path)

    with path.open("r", encoding="utf-8-sig", errors="replace", newline="") as handle:
        reader = csv.reader(handle)
        try:
            headers = next(reader)
        except StopIteration:
            return []
        symbol_idx = _find_header(headers, SYMBOL_ALIASES)
        name_idx = _find_header(headers, NAME_ALIASES)
        sector_idx = _find_header(headers, SECTOR_ALIASES)
        group_idx = _find_header(headers, GROUP_ALIASES)
        security_id_idx = _find_header(headers, SECURITY_ID_ALIASES)
        isin_idx = _find_header(headers, ISIN_ALIASES)
        if symbol_idx is None:
            raise ValueError(f"{path.name}: no symbol/security-code column found. Headers={headers}")

        records: List[UniverseRecord] = []
        for raw in reader:
            # BSE exports commonly append empty cells beyond the official header.
            row = list(raw[:len(headers)]) + [""] * max(0, len(headers) - len(raw))
            raw_symbol = _clean(row[symbol_idx])
            if not raw_symbol:
                continue

            numeric_symbol = _normalise_security_code(raw_symbol)
            is_bse = bool(re.fullmatch(r"\d{6}", numeric_symbol))
            if is_bse:
                symbol = numeric_symbol
                exchange = "BSE"
            else:
                symbol = re.sub(r"\.(BO|NS)$", "", raw_symbol.upper(), flags=re.IGNORECASE)
                symbol = re.sub(r"\s+", "", symbol)
                exchange = "NSE"
                if not symbol:
                    continue

            name = _clean(row[name_idx]) if name_idx is not None else ""
            sector = _clean(row[sector_idx]) if sector_idx is not None else ""
            group = _clean(row[group_idx]).strip().upper() if group_idx is not None else ""
            group = group[:1] if group else ""
            isin = _clean(row[isin_idx]) if isin_idx is not None else ""
            security_id_raw = _clean(row[security_id_idx]) if security_id_idx is not None else ""
            security_id = _valid_security_id(security_id_raw, symbol=symbol, isin=isin)

            # In a non-BSE CSV, a Security Id column may itself be the real symbol.
            if exchange == "NSE" and security_id:
                symbol = security_id

            ticker = _preferred_ticker(symbol, security_id, exchange)
            records.append(UniverseRecord(
                symbol=symbol,
                ticker=ticker,
                company_name=name or security_id or f"Name unavailable ({symbol})",
                exchange=exchange,
                bse_group=group,
                sector=sector,
                source_universes=[label],
                security_id=security_id,
            ))
    return records


def _merge_into(current: UniverseRecord, incoming: UniverseRecord) -> UniverseRecord:
    for source in incoming.source_universes:
        if source not in current.source_universes:
            current.source_universes.append(source)

    if current.company_name.startswith("Name unavailable") and not incoming.company_name.startswith("Name unavailable"):
        current.company_name = incoming.company_name
    if not current.sector and incoming.sector:
        current.sector = incoming.sector
    if not current.bse_group and incoming.bse_group:
        current.bse_group = incoming.bse_group
    if not current.security_id and incoming.security_id:
        current.security_id = incoming.security_id
    if current.exchange != "BSE" and incoming.exchange == "BSE":
        current.exchange = "BSE"
    current.refresh_preferred_ticker()
    return current


def load_and_deduplicate_universes(paths: Sequence[str]) -> List[UniverseRecord]:
    merged: Dict[str, UniverseRecord] = {}
    for path in paths:
        for record in load_universe_file(path):
            key = record.stable_key
            current = merged.get(key)
            if current is None:
                record.refresh_preferred_ticker()
                merged[key] = record
            else:
                merged[key] = _merge_into(current, record)
    return sorted(merged.values(), key=lambda r: (r.exchange, r.symbol))


def load_latest_bse_gainers_from_workbook(xlsx_path: str) -> List[UniverseRecord]:
    """Read the first gainer table from the latest BSEGain_* sheet.

    The gainer sheet usually contains only the six-digit code and company name.
    During the later merge step, matching Group A/B records enrich it with the
    Security Id used for Yahoo ticker resolution.
    """
    from openpyxl import load_workbook

    path = Path(xlsx_path)
    if not path.exists():
        return []
    workbook = load_workbook(path, read_only=True, data_only=True)
    sheet_names = [name for name in workbook.sheetnames if name.startswith("BSEGain_")]
    if not sheet_names:
        workbook.close()
        return []
    ws = workbook[sheet_names[-1]]
    symbol_col = name_col = sector_col = None
    records: List[UniverseRecord] = []
    located = False
    for row in ws.iter_rows(values_only=True):
        values = ["" if value is None else str(value).strip() for value in row]
        normalised = [_norm(value) for value in values]
        if not located:
            if "symbol" in normalised and ("name" in normalised or "company" in normalised):
                symbol_col = normalised.index("symbol")
                name_col = normalised.index("name") if "name" in normalised else normalised.index("company")
                sector_col = normalised.index("sector") if "sector" in normalised else None
                located = True
            continue
        symbol = _clean(values[symbol_col]) if symbol_col is not None and symbol_col < len(values) else ""
        if not symbol:
            if records:
                break
            continue
        if _norm(symbol) in {"bsesmallcapgainersdeepdive", "nextdayexplosivemovewatchlist"}:
            break
        if not re.fullmatch(r"[A-Za-z0-9&._-]+", symbol):
            break
        if re.fullmatch(r"\d+\.0+", symbol):
            symbol = symbol.split(".", 1)[0]
        name = _clean(values[name_col]) if name_col is not None and name_col < len(values) else ""
        if name and (name.isdigit() or _norm(name) == _norm(symbol)):
            name = ""
        sector = _clean(values[sector_col]) if sector_col is not None and sector_col < len(values) else ""
        exchange = "BSE" if symbol.isdigit() else "NSE"
        ticker = f"{symbol}.BO" if exchange == "BSE" else f"{symbol.upper()}.NS"
        records.append(UniverseRecord(
            symbol=symbol.upper(),
            ticker=ticker,
            company_name=name or f"Name unavailable ({symbol})",
            exchange=exchange,
            bse_group="",
            sector=sector,
            source_universes=["BSE Gainer"],
            security_id="",
        ))
    workbook.close()
    return records


def merge_universe_records(record_groups: Sequence[Sequence[UniverseRecord]]) -> List[UniverseRecord]:
    """Deduplicate parsed records while retaining sources and ticker metadata."""
    merged: Dict[str, UniverseRecord] = {}
    for records in record_groups:
        for record in records:
            key = record.stable_key
            current = merged.get(key)
            if current is None:
                record.refresh_preferred_ticker()
                merged[key] = record
            else:
                merged[key] = _merge_into(current, record)
    return sorted(merged.values(), key=lambda r: (r.exchange, r.symbol))
