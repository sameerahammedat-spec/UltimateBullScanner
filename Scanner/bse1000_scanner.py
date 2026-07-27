"""
BSE 1000 Scanner — Full Scan + Deep Dive + Ultra Top Picks.

WHY A THIRD TABLE (Ultra Top Picks): BSE smallcaps/microcaps regularly gain
14%+ in a single day on thin volume, then give it all back within a week -
that's a blow-off spike, not a real trend, and it shows up here as HIGH
today's-gain but LOW Trend Quality (ADX x R^2) because there was no base
built beforehand. Ultra Top Picks deliberately excludes those: it requires
(a) a recognizable chart-pattern setup (via chart_patterns.py - OpenCV-based
swing/trendline geometry, not just a percentage filter), (b) real Trend
Quality, (c) reasonable liquidity, and it explicitly REJECTS anything flagged
as a Blow-off Spike, no matter how big today's gain looks.

Read chart_patterns.py's module docstring before trusting the pattern label
too much - it's a heuristic geometric screen, not a crystal ball.

USAGE:
    python bse1000_scanner.py trading_toolkit.xlsx bse1000_clean.csv

(Run fix_bse_universe.py first if you haven't already, to turn your raw
BSE index export into a clean Symbol-based CSV - see GETTING_BSE_SCRIP_MASTER.md)
"""

import sys
import time
from datetime import date

from openpyxl import load_workbook
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.utils import get_column_letter

from market_scanner import (
    FONT, BLACK, HDR_FONT, HDR_FILL, SEC_FILL, SEC_FONT, ALT_FILL, BORDER,
    load_universe, analyze_symbol, analyze_universe_parallel, get_nifty_returns, score_row, verdict, TOP_N_DEEPDIVE,
)
from chart_patterns import detect_pattern, BULLISH_PATTERNS

# --- Ultra Top Picks thresholds (tune these to taste) ----------------------
ULTRA_MIN_TREND_QUALITY = 15.0     # ADX x R^2 - both strong AND clean
ULTRA_MIN_PATTERN_CONFIDENCE = 55.0
ULTRA_MIN_TURNOVER_CR = 1.0        # Rs 1 Cr/day traded value floor - avoids illiquid micro-caps
ULTRA_MAX_CANDIDATES = 20

FULL_HEADERS = ["Symbol", "Name", "Cap Segment", "Close", "% Chg", "Volume", "Turnover (Rs Cr)",
                "Liquidity OK?", "RSI(14)", "ADX(14)", "Trend Quality", "Score (0-9)", "Verdict"]

DEEPDIVE_HEADERS = ["Symbol", "Name", "Close", "Score", "Verdict", "Vol Ratio (3d)",
                     "ADX(14)", "Trend Direction", "Trend R^2", "Trend Quality (ADX x R^2)",
                     "Entry", "Stop Loss", "Target 1", "Target 2"]

ULTRA_HEADERS = ["Symbol", "Name", "Close", "% Chg", "Turnover (Rs Cr)", "Trend Quality",
                  "Chart Pattern", "Pattern Confidence", "Pattern Detail",
                  "Entry", "Stop Loss", "Target 1", "Target 2"]


def run_bse1000_scan(universe_paths, verbose=True):
    universe = load_universe(universe_paths)
    symbols = universe["symbol"].tolist()
    meta = universe.set_index("symbol").to_dict(orient="index")

    if verbose:
        print(f"BSE 1000 universe loaded: {len(symbols)} symbols")
    nifty_1m, nifty_3m = get_nifty_returns()

    results = []
    fetched = analyze_universe_parallel(symbols, nifty_1m, nifty_3m, return_series=True, verbose=verbose)
    for sym, d in fetched.items():
        if d is None or not d.get("liquidity_ok"):
            continue

        d["symbol"] = sym
        d["name"] = meta.get(sym, {}).get("name", sym)
        d["cap_segment"] = meta.get(sym, {}).get("cap_segment", "")
        d["score"] = score_row(d)
        d["verdict"] = verdict(d["score"])
        d["turnover_cr"] = (d["close"] * d["volume"]) / 1e7 if (d["close"] and d["volume"]) else None

        # Pattern detection is pure CPU/numpy work (no network), so it runs
        # fast in the main thread after all the network fetching is done -
        # no benefit to parallelizing this part, and cv2 objects aren't
        # guaranteed thread-safe to share across worker threads anyway.
        pattern_result = detect_pattern(d["close_series"], d["volume_series"])
        d["pattern"] = pattern_result["pattern"]
        d["pattern_confidence"] = pattern_result["confidence"]
        d["pattern_detail"] = pattern_result["detail"]
        d["is_spike_warning"] = pattern_result["is_spike_warning"]

        results.append(d)

    full_scan = sorted(results, key=lambda r: r["score"], reverse=True)
    deep_dive = [d for d in full_scan if d["verdict"] in ("STRONG SETUP", "WATCHLIST")][:TOP_N_DEEPDIVE]

    ultra = [
        d for d in full_scan
        if not d["is_spike_warning"]
        and d.get("trend_quality") is not None and d["trend_quality"] >= ULTRA_MIN_TREND_QUALITY
        and d["pattern"] in BULLISH_PATTERNS
        and d["pattern_confidence"] >= ULTRA_MIN_PATTERN_CONFIDENCE
        and d.get("turnover_cr") is not None and d["turnover_cr"] >= ULTRA_MIN_TURNOVER_CR
    ]
    ultra = sorted(ultra, key=lambda d: d["pattern_confidence"], reverse=True)[:ULTRA_MAX_CANDIDATES]

    return full_scan, deep_dive, ultra


# --- Excel writers -----------------------------------------------------

def _write_section(ws, start_row, headers, rows, row_builder, title=None):
    r = start_row
    if title:
        ws.cell(row=r, column=1, value=title).font = SEC_FONT
        for c in range(1, len(headers) + 1):
            ws.cell(row=r, column=c).fill = SEC_FILL
        ws.merge_cells(start_row=r, start_column=1, end_row=r, end_column=len(headers))
        r += 1
    for i, h in enumerate(headers, start=1):
        cell = ws.cell(row=r, column=i, value=h)
        cell.font = HDR_FONT
        cell.fill = HDR_FILL
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        cell.border = BORDER
    ws.row_dimensions[r].height = 32
    r += 1
    for idx, d in enumerate(rows):
        vals = row_builder(d)
        for c, v in enumerate(vals, start=1):
            cell = ws.cell(row=r, column=c, value=v)
            cell.font = BLACK
            cell.border = BORDER
            cell.alignment = Alignment(horizontal="center", wrap_text=True)
            if idx % 2 == 1:
                cell.fill = ALT_FILL
        r += 1
    return r + 1


def num(v, nd=2):
    return round(v, nd) if v is not None else "N/A"


def full_row(d):
    return [d["symbol"], d["name"], d.get("cap_segment", ""), num(d["close"]), num(d["pct_change"]),
            d["volume"], num(d.get("turnover_cr")), "Y" if d["liquidity_ok"] else "N",
            num(d.get("rsi"), 1), num(d.get("adx"), 1), num(d.get("trend_quality"), 1),
            d["score"], d["verdict"]]


def deepdive_row(d):
    return [d["symbol"], d["name"], num(d["close"]), d["score"], d["verdict"], num(d.get("vol_ratio_3d")),
            num(d.get("adx"), 1), d.get("trend_direction", "N/A"), num(d.get("trend_r2"), 2),
            num(d.get("trend_quality"), 1), num(d.get("entry")), num(d.get("stop_loss")),
            num(d.get("target1")), num(d.get("target2"))]


def ultra_row(d):
    return [d["symbol"], d["name"], num(d["close"]), num(d["pct_change"]), num(d.get("turnover_cr")),
            num(d.get("trend_quality"), 1), d["pattern"], num(d["pattern_confidence"], 1),
            d["pattern_detail"], num(d.get("entry")), num(d.get("stop_loss")),
            num(d.get("target1")), num(d.get("target2"))]


def main(xlsx_path, universe_paths):
    full_scan, deep_dive, ultra = run_bse1000_scan(universe_paths)
    today = date.today()
    sheet_name = f"BSE1000_{today.strftime('%d%b%y')}"

    wb = load_workbook(xlsx_path)
    if sheet_name in wb.sheetnames:
        del wb[sheet_name]
    ws = wb.create_sheet(sheet_name)
    ws.sheet_view.showGridLines = False
    ws.cell(row=1, column=1, value=f"BSE 1000 Scan — {today.strftime('%d %b %Y')}").font = \
        Font(name=FONT, size=14, bold=True, color="1F3864")
    ws.merge_cells(start_row=1, start_column=1, end_row=1, end_column=len(FULL_HEADERS))

    r = _write_section(ws, 3, FULL_HEADERS, full_scan, full_row,
                        title=f"FULL SCAN — {len(full_scan)} liquid stocks scanned")
    r = _write_section(ws, r, DEEPDIVE_HEADERS, deep_dive, deepdive_row,
                        title=f"DEEP DIVE — top {len(deep_dive)} by score")
    _write_section(ws, r, ULTRA_HEADERS, ultra, ultra_row,
                   title=f"ULTRA TOP PICKS — {len(ultra)} (pattern-confirmed, spike-filtered, "
                         f"trend quality >= {ULTRA_MIN_TREND_QUALITY})")

    for i in range(1, len(FULL_HEADERS) + 1):
        ws.column_dimensions[get_column_letter(i)].width = 15
    ws.column_dimensions["B"].width = 26
    wb.save(xlsx_path)

    print(f"\nDone. Sheet '{sheet_name}': {len(full_scan)} scanned, {len(deep_dive)} deep-dive, "
          f"{len(ultra)} Ultra Top Picks.")
    print("Ultra Top Picks excludes anything geometrically flagged as a blow-off spike, "
          "regardless of how large today's % gain is - that's by design, re-read the "
          "chart_patterns.py docstring if a pick you expected is missing.")


if __name__ == "__main__":
    if len(sys.argv) < 3:
        print("Usage: python bse1000_scanner.py trading_toolkit.xlsx bse1000_clean.csv [more.csv ...]")
        sys.exit(1)
    main(sys.argv[1], sys.argv[2:])
