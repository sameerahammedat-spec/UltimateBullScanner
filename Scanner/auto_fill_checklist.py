"""
Auto-fill technical columns in trading_toolkit.xlsx (Scanner Shortlist sheet)
using free Yahoo Finance data via yfinance.

Fills automatically:  H (Above 50 DMA), I (Above 200 DMA), J (RSI-14), L (Volume Spike)
Leaves manual (by design — need judgment/paid data): M, N, O

Setup (once):
    pip install yfinance openpyxl pandas

Usage:
    python auto_fill_checklist.py trading_toolkit.xlsx
"""

import sys
import time
import pandas as pd
import yfinance as yf
from openpyxl import load_workbook


def rsi(series, period=14):
    delta = series.diff()
    gain = delta.clip(lower=0)
    loss = -delta.clip(upper=0)
    avg_gain = gain.rolling(period).mean()
    avg_loss = loss.rolling(period).mean()
    rs = avg_gain / avg_loss
    return 100 - (100 / (1 + rs))


def to_scalar(x):
    """Safely pull a single float out of anything pandas hands back
    (a plain float, a 1-item Series, or a 1-item DataFrame column)."""
    if isinstance(x, pd.Series):
        x = x.iloc[-1] if len(x) else None
    if x is None or pd.isna(x):
        return None
    return float(x)


def analyze_symbol(nse_symbol):
    """Fetch ~1y of daily data and compute the 4 automatable checks."""
    ticker_str = nse_symbol.strip() + ".NS"   # NSE suffix for Yahoo Finance
    try:
        tk = yf.Ticker(ticker_str)
        df = tk.history(period="1y", interval="1d", auto_adjust=True)
        if df.empty or len(df) < 60:
            return None
    except Exception as e:
        print(f"  [!] {nse_symbol}: fetch failed ({e})")
        return None

    # Safety net: flatten in the rare case columns come back multi-level
    if isinstance(df.columns, pd.MultiIndex):
        df.columns = df.columns.get_level_values(0)

    close = df["Close"]
    volume = df["Volume"]

    last_close = to_scalar(close.iloc[-1])
    dma50 = to_scalar(close.rolling(50).mean().iloc[-1])
    dma200 = to_scalar(close.rolling(200).mean().iloc[-1]) if len(close) >= 200 else None
    rsi14 = to_scalar(rsi(close, 14).iloc[-1])

    avg_vol_20 = to_scalar(volume.rolling(20).mean().iloc[-1])
    last_vol = to_scalar(volume.iloc[-1])
    vol_spike = None
    if avg_vol_20 is not None and last_vol is not None and avg_vol_20 > 0:
        vol_spike = last_vol > 1.5 * avg_vol_20

    return {
        "above_50dma": "Y" if (dma50 is not None and last_close is not None and last_close > dma50) else "N",
        "above_200dma": ("Y" if (dma200 is not None and last_close is not None and last_close > dma200) else "N") if dma200 is not None else "",
        "rsi": round(rsi14, 1) if rsi14 is not None else "",
        "vol_spike": "Y" if vol_spike else ("N" if vol_spike is not None else ""),
    }


def main(filepath):
    wb = load_workbook(filepath)
    ws = wb["Scanner Shortlist"]

    header_row = 4
    data_start = header_row + 1
    last_row = ws.max_row

    print(f"Scanning rows {data_start} to {last_row}...\n")

    for r in range(data_start, last_row + 1):
        symbol = ws.cell(row=r, column=3).value   # column C = Symbol
        name = ws.cell(row=r, column=2).value
        if not symbol:
            continue

        print(f"[{r - data_start + 1}] {symbol} ({name})...", end=" ", flush=True)

        try:
            result = analyze_symbol(symbol)
        except Exception as e:
            print(f"SKIPPED (unexpected error: {e})")
            continue

        if result is None:
            print("SKIPPED (no data / delisted / illiquid on Yahoo)")
            continue

        ws.cell(row=r, column=8, value=result["above_50dma"])       # H
        if result["above_200dma"] != "":
            ws.cell(row=r, column=9, value=result["above_200dma"])  # I
        if result["rsi"] != "":
            ws.cell(row=r, column=10, value=result["rsi"])          # J
        if result["vol_spike"] != "":
            ws.cell(row=r, column=12, value=result["vol_spike"])    # L

        print(f"50DMA={result['above_50dma']} 200DMA={result['above_200dma']} "
              f"RSI={result['rsi']} VolSpike={result['vol_spike']}")

        time.sleep(0.3)  # be polite to Yahoo's servers

    wb.save(filepath)
    print(f"\nDone. Saved back to {filepath}")
    print("Columns M (Support/Breakout), N (FII/DII), O (Fundamentals) still need your manual check.")


if __name__ == "__main__":
    if len(sys.argv) != 2:
        print("Usage: python auto_fill_checklist.py trading_toolkit.xlsx")
        sys.exit(1)
    main(sys.argv[1])