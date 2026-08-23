"""BSE 2:00 PM-3:10 PM Fresh Big-Move Scanner.

Purpose
-------
Scan the union of BSE Group A, Group B and BSE Smallcap constituents for a move
that *started after 2 PM*, validate it with same-time volume, daily technical
structure, historical 3-session continuation behaviour and fresh catalysts,
then optionally alert by SMS/voice/email.

This is deliberately separate from the existing end-of-day scanners. No
existing workbook, scanner or Explosive Engine code path is changed.

Examples
--------
    python afternoon_fresh_move_scanner.py --once
    python afternoon_fresh_move_scanner.py --watch
"""
from __future__ import annotations

import argparse
import json
import math
import os
import sys
import time as time_module
from dataclasses import asdict
from datetime import date, datetime, timedelta, time
from pathlib import Path
from typing import Any, Dict, Iterable, List, Optional, Sequence, Tuple
from zoneinfo import ZoneInfo

import numpy as np
import pandas as pd
import requests
import yfinance as yf
from openpyxl import load_workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from afternoon_alerts import send_candidate_alerts
from afternoon_catalyst import CatalystAssessment, assess_catalyst, fetch_bse_announcements
from afternoon_config import AfternoonConfig, DEFAULT_AFTERNOON_CONFIG
from afternoon_momentum_engine import (
    PrefilterResult,
    build_risk_levels,
    calibrate_historical_continuation,
    clean_ohlcv,
    compute_daily_features,
    compute_intraday_prefilter,
    compute_intraday_rsi,
    compute_same_window_rvol,
    score_candidate,
)
from market_scanner import _SHARED_SESSION


DEFAULT_GROUP_A = "Group_A.csv"
DEFAULT_GROUP_B = "Group_B.csv"
DEFAULT_SMALLCAP = "bse_smallcap250_clean.csv"
DEFAULT_OUTPUT_DIR = "afternoon_reports"
BSE_GAINER_URL = "https://api.bseindia.com/BseIndiaAPI/api/MktRGainerLoserData/w"

LATEST_COLUMNS = [
    "scan_time", "rank", "security_code", "security_id", "name", "group", "is_smallcap",
    "stage0_source", "stage0_day_change_pct", "latest_price", "window_move_pct", "pre_window_move_pct", "ret_15m_pct", "ret_30m_pct",
    "window_turnover_cr", "same_window_rvol", "volume_acceleration", "distance_from_window_high_pct",
    "intraday_rsi", "daily_rsi", "atr_pct", "prior_2d_return_pct", "prior_5d_return_pct",
    "breakout_20d_pct", "room_to_resistance_pct", "market_window_move_pct", "score", "status",
    "entry", "stop_loss", "risk_pct", "target1", "target2", "rr_target1", "rr_target2",
    "hist_samples", "hist_hit_5_pct", "hist_hit_8_pct", "hist_hit_10_pct", "hist_hit_15_pct", "hist_hit_20_pct",
    "hist_p75_mfe_pct", "catalyst_strength", "catalyst_reason", "catalyst_evidence",
    "reasons", "warnings", "hard_rejects",
]


# ---------------------------------------------------------------------------
# Universe loading
# ---------------------------------------------------------------------------

def _read_csv_robust(path: str) -> pd.DataFrame:
    # Group_A.csv/Group_B.csv in the supplied project have harmless trailing
    # commas after the declared columns. index_col=False prevents pandas from
    # accidentally treating the first data columns as an inferred multi-index.
    try:
        return pd.read_csv(path, dtype=str, keep_default_na=False, index_col=False)
    except UnicodeDecodeError:
        return pd.read_csv(path, dtype=str, keep_default_na=False, index_col=False, encoding="latin-1")


def _clean(value: Any) -> str:
    text = str(value or "").strip()
    return "" if text.lower() in ("nan", "none", "null") else text


def _valid_bse_id(value: str) -> bool:
    return bool(value and value not in ("-", "--") and len(value) <= 30)


def load_afternoon_universe(group_a_path: str, group_b_path: str, smallcap_path: str) -> List[Dict[str, Any]]:
    smallcap_df = _read_csv_robust(smallcap_path)
    smallcap_isins = set(_clean(x).upper() for x in smallcap_df.get("ISIN", pd.Series(dtype=str)) if _clean(x))
    smallcap_symbols = set(_clean(x).upper() for x in smallcap_df.get("Symbol", pd.Series(dtype=str)) if _clean(x))

    records: Dict[str, Dict[str, Any]] = {}
    for path, group_label in ((group_a_path, "A"), (group_b_path, "B")):
        df = _read_csv_robust(path)
        required = ["Security Code", "Security Id"]
        missing = [c for c in required if c not in df.columns]
        if missing:
            raise ValueError(f"{path}: missing required columns {missing}; found {list(df.columns)}")
        for _, row in df.iterrows():
            security_code = _clean(row.get("Security Code"))
            security_id = _clean(row.get("Security Id")).upper()
            isin = _clean(row.get("ISIN No")).upper()
            status = _clean(row.get("Status")).lower()
            if status and status != "active":
                continue
            if not security_id and not security_code:
                continue
            name = _clean(row.get("Security Name")) or _clean(row.get("Issuer Name")) or security_id or security_code
            key = isin or f"{group_label}:{security_code or security_id}"
            primary = f"{security_id}.BO" if _valid_bse_id(security_id) else f"{security_code}.BO"
            alternates = []
            if security_code:
                alternates.append(f"{security_code}.BO")
            if security_id:
                alternates.append(f"{security_id}.NS")
            record = {
                "security_code": security_code,
                "security_id": security_id,
                "name": name,
                "group": group_label,
                "isin": isin,
                "is_smallcap": bool((isin and isin in smallcap_isins) or security_id in smallcap_symbols),
                "ticker": primary,
                "alternate_tickers": list(dict.fromkeys(x for x in alternates if x != primary)),
            }
            records[key] = record

    # Any smallcap constituent missing from A/B is still included so the stated
    # universe remains BSE Smallcap + A/B, not merely their current overlap.
    known_isins = {r.get("isin") for r in records.values() if r.get("isin")}
    for _, row in smallcap_df.iterrows():
        isin = _clean(row.get("ISIN")).upper()
        symbol = _clean(row.get("Symbol")).upper()
        if not symbol or (isin and isin in known_isins):
            continue
        name = _clean(row.get("Company Name")) or symbol
        key = isin or f"SMALLCAP:{symbol}"
        records[key] = {
            "security_code": "",
            "security_id": symbol,
            "name": name,
            "group": "",
            "isin": isin,
            "is_smallcap": True,
            "ticker": f"{symbol}.BO",
            "alternate_tickers": [f"{symbol}.NS"],
        }

    result = list(records.values())
    result.sort(key=lambda r: (r.get("group") != "A", not r.get("is_smallcap"), r.get("security_id", "")))
    return result


# ---------------------------------------------------------------------------
# BSE stage-0 live mover shortlist
# ---------------------------------------------------------------------------

def _first_value(record: Dict[str, Any], keys: Sequence[str]) -> Any:
    normalized = {str(k).strip().lower().replace("_", "").replace(" ", ""): v for k, v in record.items()}
    for key in keys:
        norm = key.strip().lower().replace("_", "").replace(" ", "")
        if norm in normalized and str(normalized[norm]).strip() not in ("", "None", "nan"):
            return normalized[norm]
    return None


def _float_value(record: Dict[str, Any], keys: Sequence[str]) -> Optional[float]:
    value = _first_value(record, keys)
    if value is None:
        return None
    try:
        return float(str(value).replace(",", "").replace("%", "").strip())
    except (TypeError, ValueError):
        return None


def fetch_bse_live_mover_shortlist(
    universe: Sequence[Dict[str, Any]],
    config: AfternoonConfig = DEFAULT_AFTERNOON_CONFIG,
) -> Tuple[Optional[List[Dict[str, Any]]], str]:
    """Best-effort BSE live shortlist.

    The endpoint is used only to reduce the number of symbols sent to Yahoo for
    5-minute verification. A stock NEVER qualifies on this snapshot alone.
    Returning None means the source failed and the caller should use fallback
    coverage; returning [] means BSE responded successfully but no universe
    members met the light stage-0 gate.
    """
    if not config.stage0_bse_enabled:
        return None, "BSE_STAGE0_DISABLED"

    by_code = {str(x.get("security_code") or "").strip(): x for x in universe if x.get("security_code")}
    by_id = {str(x.get("security_id") or "").strip().upper(): x for x in universe if x.get("security_id")}
    headers = {
        "User-Agent": (
            "Mozilla/5.0 (X11; Linux x86_64) AppleWebKit/537.36 "
            "(KHTML, like Gecko) Chrome/145.0 Safari/537.36"
        ),
        "Referer": "https://www.bseindia.com/markets/equity/EQReports/MarketWatch.html",
        "Accept": "application/json,text/plain,*/*",
    }
    session = requests.Session()
    selected: Dict[str, Dict[str, Any]] = {}
    successful_calls = 0

    # Percentage, traded value and traded volume orderings complement each
    # other. This helps retain a newly accelerating stock whose full-day gain is
    # still modest but whose afternoon participation is already exceptional.
    for order_by in ("all", "value", "volume"):
        params = {
            "GLtype": "gainer",
            "IndxGrp": "AllMkt",
            "IndxGrpval": "AllMkt",
            "orderby": order_by,
        }
        try:
            response = session.get(
                BSE_GAINER_URL,
                headers=headers,
                params=params,
                timeout=config.request_timeout_seconds,
            )
            response.raise_for_status()
            payload = response.json()
        except Exception as exc:
            print(f"[Afternoon stage0] BSE {order_by} snapshot unavailable: {exc}")
            continue

        rows = payload.get("Table") or payload.get("table") or payload.get("data") or []
        if not isinstance(rows, list):
            continue
        successful_calls += 1
        retained_this_order = 0
        for row in rows:
            if not isinstance(row, dict):
                continue
            code = str(_first_value(row, ["scrip_code", "SCRIP_CD", "ScripCode", "Security Code"]) or "").strip()
            sid = str(_first_value(row, ["scrip_name", "SCRIP_NAME", "SecurityId", "symbol"]) or "").strip().upper()
            stock = by_code.get(code) or by_id.get(sid)
            if stock is None:
                continue
            day_change = _float_value(
                row,
                ["change_percent", "percent_change", "chgper", "perchg", "change", "% change", "pcntChange"],
            )
            # The endpoint itself is gainers-only. Do not hard-filter on the
            # response's change field because BSE payload field names/units have
            # changed over time; strict freshness is verified from 5-minute bars.
            key = str(stock.get("security_code") or stock.get("security_id") or stock.get("ticker"))
            item = dict(stock)
            item["stage0_source"] = f"BSE_{order_by.upper()}"
            item["stage0_day_change_pct"] = day_change
            if key not in selected:
                selected[key] = item
                retained_this_order += 1
            if retained_this_order >= config.stage0_each_order_limit:
                break

    if successful_calls == 0:
        return None, "BSE_STAGE0_FAILED"
    shortlist = list(selected.values())
    shortlist.sort(
        key=lambda x: (
            x.get("group") != "A",
            not x.get("is_smallcap", False),
            -(x.get("stage0_day_change_pct") or 0.0),
        )
    )
    return shortlist, f"BSE_STAGE0_{successful_calls}_CALLS"


def yahoo_fallback_universe(
    universe: Sequence[Dict[str, Any]],
    cycle_number: int,
    config: AfternoonConfig = DEFAULT_AFTERNOON_CONFIG,
) -> List[Dict[str, Any]]:
    """Controlled fallback when BSE's stage-0 webpage endpoint is unavailable.

    Group A and all tagged Smallcap names are always covered. Group B is
    deterministically sharded across cycles to avoid hammering Yahoo with 2,300
    intraday histories every five minutes. In production, a broker/exchange
    websocket should replace this fallback for guaranteed full-universe latency.
    """
    shards = max(1, int(config.fallback_group_b_shards))
    shard = cycle_number % shards
    result = []
    seen = set()
    for stock in universe:
        include = stock.get("group") == "A" or bool(stock.get("is_smallcap"))
        if not include and stock.get("group") == "B":
            raw = str(stock.get("security_code") or stock.get("security_id") or "0")
            try:
                bucket = int("".join(ch for ch in raw if ch.isdigit()) or "0") % shards
            except ValueError:
                bucket = sum(ord(ch) for ch in raw) % shards
            include = bucket == shard
        key = stock.get("security_code") or stock.get("security_id") or stock.get("ticker")
        if include and key not in seen:
            seen.add(key)
            item = dict(stock)
            item["stage0_source"] = f"YAHOO_FALLBACK_B_SHARD_{shard + 1}_OF_{shards}"
            result.append(item)
    return result


# ---------------------------------------------------------------------------
# Yahoo batch adapter
# ---------------------------------------------------------------------------

def _chunks(items: Sequence[Any], size: int) -> Iterable[Sequence[Any]]:
    for i in range(0, len(items), size):
        yield items[i : i + size]


def _extract_ticker_frame(downloaded: pd.DataFrame, ticker: str, ticker_count: int) -> pd.DataFrame:
    if downloaded is None or downloaded.empty:
        return pd.DataFrame()
    if not isinstance(downloaded.columns, pd.MultiIndex):
        return downloaded.copy() if ticker_count == 1 else pd.DataFrame()

    level0 = set(str(x) for x in downloaded.columns.get_level_values(0))
    level1 = set(str(x) for x in downloaded.columns.get_level_values(1)) if downloaded.columns.nlevels >= 2 else set()
    try:
        if ticker in level0:
            frame = downloaded[ticker].copy()
        elif ticker in level1:
            frame = downloaded.xs(ticker, axis=1, level=1).copy()
        else:
            return pd.DataFrame()
    except Exception:
        return pd.DataFrame()
    if isinstance(frame, pd.Series):
        frame = frame.to_frame()
    return frame


def batch_download(
    tickers: Sequence[str],
    period: str,
    interval: str,
    batch_size: int,
    auto_adjust: bool = False,
) -> Dict[str, pd.DataFrame]:
    results: Dict[str, pd.DataFrame] = {}
    unique = list(dict.fromkeys(t for t in tickers if t))
    for batch_no, batch in enumerate(_chunks(unique, batch_size), start=1):
        print(f"[Afternoon] Yahoo {interval}/{period}: batch {batch_no}, {len(batch)} tickers")
        try:
            downloaded = yf.download(
                tickers=list(batch),
                period=period,
                interval=interval,
                group_by="ticker",
                auto_adjust=auto_adjust,
                progress=False,
                threads=True,
                prepost=False,
                timeout=20,
                session=_SHARED_SESSION,
            )
        except Exception as exc:
            # Newer yfinance versions use curl_cffi and may reject a
            # requests.Session; older versions used it successfully. Retry once
            # without a custom session for forward/backward compatibility.
            message = str(exc).lower()
            if isinstance(exc, TypeError) or "session" in message or "curl_cffi" in message:
                try:
                    downloaded = yf.download(
                        tickers=list(batch), period=period, interval=interval,
                        group_by="ticker", auto_adjust=auto_adjust, progress=False,
                        threads=True, prepost=False, timeout=20,
                    )
                except Exception as retry_exc:
                    print(f"[Afternoon] batch download failed after sessionless retry: {retry_exc}")
                    continue
            else:
                print(f"[Afternoon] batch download failed: {exc}")
                continue
        for ticker in batch:
            frame = _extract_ticker_frame(downloaded, ticker, len(batch))
            if frame is not None and not frame.empty:
                results[ticker] = frame
    return results


def _fallback_one(stock: Dict[str, Any], period: str, interval: str) -> tuple[str, pd.DataFrame]:
    choices = [stock.get("ticker"), *stock.get("alternate_tickers", [])]
    for ticker_symbol in dict.fromkeys(x for x in choices if x):
        try:
            try:
                ticker_obj = yf.Ticker(ticker_symbol, session=_SHARED_SESSION)
                frame = ticker_obj.history(period=period, interval=interval, auto_adjust=False)
            except Exception as session_exc:
                if "session" not in str(session_exc).lower() and "curl_cffi" not in str(session_exc).lower():
                    raise
                frame = yf.Ticker(ticker_symbol).history(period=period, interval=interval, auto_adjust=False)
            if frame is not None and not frame.empty:
                return ticker_symbol, frame
        except Exception:
            continue
    return stock.get("ticker", ""), pd.DataFrame()


# ---------------------------------------------------------------------------
# Scanner
# ---------------------------------------------------------------------------

def _preliminary_rank(prefilter: PrefilterResult) -> float:
    move = min(prefilter.window_move_pct or 0, 7.0)
    r15 = max(0.0, prefilter.ret_15m_pct or 0.0)
    vol = min(prefilter.volume_acceleration or 1.0, 4.0)
    turnover = math.log1p(max(0.0, prefilter.window_turnover_cr or 0.0))
    near_high = max(0.0, 1.5 - (prefilter.distance_from_window_high_pct or 1.5))
    return move * 4.0 + r15 * 2.0 + vol * 3.0 + turnover * 2.0 + near_high * 2.0


def _current_clock(now: datetime, cfg: AfternoonConfig) -> time:
    return now.astimezone(ZoneInfo(cfg.timezone)).time().replace(tzinfo=None)


def _status(score: float, rejected: bool, cfg: AfternoonConfig) -> str:
    if rejected:
        return "REJECT"
    if score >= cfg.high_confidence_score:
        return "HIGH-CONFIDENCE REVIEW"
    if score >= cfg.sms_alert_score:
        return "IMMEDIATE REVIEW"
    if score >= cfg.min_report_score:
        return "WATCH"
    return "LOW"


def _flatten_candidate(candidate: Dict[str, Any]) -> Dict[str, Any]:
    row = {column: candidate.get(column) for column in LATEST_COLUMNS}
    return row


class AfternoonScanner:
    def __init__(
        self,
        universe: List[Dict[str, Any]],
        output_dir: str,
        config: AfternoonConfig = DEFAULT_AFTERNOON_CONFIG,
    ):
        self.universe = universe
        self.config = config
        self.output_dir = Path(output_dir)
        self.output_dir.mkdir(parents=True, exist_ok=True)
        self.all_snapshots: List[Dict[str, Any]] = []
        self.alert_log: List[Dict[str, Any]] = []
        self.alert_state: Dict[str, Dict[str, Any]] = {}
        self._announcements: List[Dict[str, Any]] = []
        self._announcements_fetched_at: Optional[datetime] = None
        self._cycle_number: int = 0

    def _get_announcements(self, now: datetime) -> List[Dict[str, Any]]:
        if (
            self._announcements_fetched_at is None
            or now - self._announcements_fetched_at >= timedelta(minutes=15)
        ):
            self._announcements = fetch_bse_announcements(now.date(), self.config)
            self._announcements_fetched_at = now
            print(f"[Afternoon catalyst] cached {len(self._announcements)} BSE announcements")
        return self._announcements

    def broad_scan(self, now: datetime) -> List[Dict[str, Any]]:
        self._cycle_number += 1
        stage0, stage0_status = fetch_bse_live_mover_shortlist(self.universe, self.config)
        if stage0 is None:
            scan_universe = yahoo_fallback_universe(self.universe, self._cycle_number - 1, self.config)
            source_note = f"{stage0_status}; controlled Yahoo fallback"
        else:
            scan_universe = stage0
            source_note = stage0_status

        print(
            f"[Afternoon stage0] {source_note}: verifying {len(scan_universe)} of "
            f"{len(self.universe)} universe names with 5-minute bars"
        )
        if not scan_universe:
            return []

        tickers = [stock["ticker"] for stock in scan_universe]
        frames = batch_download(
            tickers,
            period="1d",
            interval="5m",
            batch_size=self.config.broad_batch_size,
            auto_adjust=False,
        )
        candidates = []
        usable = 0
        for stock in scan_universe:
            frame = frames.get(stock["ticker"] )
            if frame is None or frame.empty:
                continue
            usable += 1
            pre = compute_intraday_prefilter(frame, now=now, config=self.config)
            if not pre.qualified:
                continue
            row = dict(stock)
            row["prefilter"] = pre
            row["preliminary_rank"] = _preliminary_rank(pre)
            candidates.append(row)

        candidates.sort(key=lambda x: x["preliminary_rank"], reverse=True)
        print(
            f"[Afternoon] verified sweep: {len(scan_universe)} stage0 | {usable} usable | "
            f"{len(candidates)} genuinely fresh post-2PM candidates"
        )
        return candidates[: self.config.deep_candidate_limit]

    def deep_scan(self, broad_candidates: List[Dict[str, Any]], now: datetime) -> Tuple[List[Dict[str, Any]], List[Dict[str, Any]]]:
        if not broad_candidates:
            return [], []
        tickers = [x["ticker"] for x in broad_candidates]
        intraday_frames = batch_download(
            tickers,
            period=self.config.intraday_history_period,
            interval="5m",
            batch_size=min(40, self.config.broad_batch_size),
            auto_adjust=False,
        )
        daily_frames = batch_download(
            tickers,
            period=self.config.daily_history_period,
            interval="1d",
            batch_size=min(60, self.config.broad_batch_size),
            auto_adjust=True,
        )

        # Sensex is only used for relative afternoon strength. Failure is a
        # missing feature, not a scanner failure.
        market_frame = batch_download(["^BSESN"], "1d", "5m", 1, auto_adjust=False).get("^BSESN", pd.DataFrame())
        market_pre = compute_intraday_prefilter(market_frame, now=now, config=self.config) if not market_frame.empty else None
        market_move = market_pre.window_move_pct if market_pre else None

        provisional: List[Dict[str, Any]] = []
        rejected: List[Dict[str, Any]] = []
        clock = _current_clock(now, self.config)
        for stock in broad_candidates:
            ticker = stock["ticker"]
            intraday = intraday_frames.get(ticker, pd.DataFrame())
            daily = daily_frames.get(ticker, pd.DataFrame())
            if intraday.empty or daily.empty:
                # Only a small shortlist reaches this point, so alternate ticker
                # fallbacks are affordable here and improve BSE symbol coverage.
                resolved, fallback_intraday = _fallback_one(stock, self.config.intraday_history_period, "5m")
                if not fallback_intraday.empty:
                    ticker = resolved
                    intraday = fallback_intraday
                if daily.empty:
                    try:
                        try:
                            daily = yf.Ticker(ticker, session=_SHARED_SESSION).history(
                                period=self.config.daily_history_period, interval="1d", auto_adjust=True
                            )
                        except Exception as session_exc:
                            if "session" not in str(session_exc).lower() and "curl_cffi" not in str(session_exc).lower():
                                raise
                            daily = yf.Ticker(ticker).history(
                                period=self.config.daily_history_period, interval="1d", auto_adjust=True
                            )
                    except Exception:
                        daily = pd.DataFrame()
            if intraday.empty or daily.empty:
                continue

            # Recompute the live window from the richer 15d frame so all deep
            # metrics reference exactly the same timestamps/data source.
            pre = compute_intraday_prefilter(intraday, now=now, config=self.config)
            if not pre.qualified:
                continue
            rvol = compute_same_window_rvol(intraday, now.date(), clock, self.config)
            i_rsi = compute_intraday_rsi(intraday, self.config)
            daily_features = compute_daily_features(
                daily,
                current_price=float(pre.latest_price),
                current_day_high=pre.day_high,
                current_day_low=pre.day_low,
                current_date=now.date(),
                config=self.config,
            )
            continuation = calibrate_historical_continuation(daily, now.date(), self.config)
            score_data = score_candidate(
                pre, rvol, i_rsi, daily_features, continuation,
                market_window_move_pct=market_move,
                catalyst_score=0.0,
                catalyst_strength="NONE",
                config=self.config,
            )
            risk = build_risk_levels(float(pre.latest_price), pre, daily_features, self.config)
            if risk.get("risk_pct") is not None and risk["risk_pct"] > self.config.max_entry_risk_pct:
                score_data["rejected"] = True
                score_data["hard_rejects"].append(
                    f"Technical invalidation requires {risk['risk_pct']:.2f}% risk (> {self.config.max_entry_risk_pct:.2f}%)"
                )

            item = {
                **stock,
                "ticker": ticker,
                **pre.to_dict(),
                "same_window_rvol": rvol,
                "intraday_rsi": i_rsi,
                "daily_rsi": daily_features.rsi14,
                **daily_features.to_dict(),
                "market_window_move_pct": market_move,
                "hist_samples": continuation.sample_size,
                "hist_hit_5_pct": continuation.hit_5_pct,
                "hist_hit_8_pct": continuation.hit_8_pct,
                "hist_hit_10_pct": continuation.hit_10_pct,
                "hist_hit_15_pct": continuation.hit_15_pct,
                "hist_hit_20_pct": continuation.hit_20_pct,
                "hist_p75_mfe_pct": continuation.p75_mfe_pct,
                **risk,
                **score_data,
                "_daily_frame": daily,
            }
            (rejected if item["rejected"] else provisional).append(item)

        provisional.sort(key=lambda x: x["score"], reverse=True)

        # Catalyst/network work only for the strongest provisional names.
        announcements = self._get_announcements(now)
        for item in provisional[: self.config.catalyst_candidate_limit]:
            catalyst = assess_catalyst(item, item["ticker"], announcements, self.config)
            rescored = score_candidate(
                PrefilterResult(**{k: item.get(k) for k in PrefilterResult.__dataclass_fields__}),
                item.get("same_window_rvol"),
                item.get("intraday_rsi"),
                compute_daily_features(
                    item["_daily_frame"], item["latest_price"], item.get("day_high"), item.get("day_low"), now.date(), self.config
                ),
                calibrate_historical_continuation(item["_daily_frame"], now.date(), self.config),
                item.get("market_window_move_pct"),
                catalyst_score=catalyst.score,
                catalyst_strength=catalyst.strength,
                config=self.config,
            )
            item.update(rescored)
            item["catalyst_strength"] = catalyst.strength
            item["catalyst_reason"] = catalyst.reason
            item["catalyst_evidence"] = catalyst.evidence
            item["catalyst_warnings"] = catalyst.warnings
            item["revenue_growth_pct"] = catalyst.revenue_growth_pct
            item["profit_growth_pct"] = catalyst.profit_growth_pct
            if catalyst.strength == "CAUTION":
                item["rejected"] = True
                item["hard_rejects"].append("Adverse/clarification catalyst detected")
            if catalyst.warnings:
                item["warnings"].append(catalyst.warnings)

        for item in provisional[self.config.catalyst_candidate_limit :]:
            item["catalyst_strength"] = "NOT_CHECKED"
            item["catalyst_reason"] = "Catalyst check reserved for top-ranked candidates"
            item["catalyst_evidence"] = ""

        qualified = []
        for item in provisional:
            item["status"] = _status(item["score"], item["rejected"], self.config)
            if item["rejected"]:
                rejected.append(item)
            elif item["score"] >= self.config.min_report_score:
                qualified.append(item)

        qualified.sort(key=lambda x: x["score"], reverse=True)
        rejected.sort(key=lambda x: x.get("score", 0), reverse=True)
        for item in qualified + rejected:
            item.pop("_daily_frame", None)
            item["reasons"] = "; ".join(item.get("reasons", [])) if isinstance(item.get("reasons"), list) else item.get("reasons", "")
            item["warnings"] = "; ".join(str(x) for x in item.get("warnings", []) if x) if isinstance(item.get("warnings"), list) else item.get("warnings", "")
            item["hard_rejects"] = "; ".join(item.get("hard_rejects", [])) if isinstance(item.get("hard_rejects"), list) else item.get("hard_rejects", "")
        return qualified, rejected

    def _should_alert(self, item: Dict[str, Any], now: datetime) -> bool:
        if item.get("rejected") or float(item.get("score") or 0) < self.config.sms_alert_score:
            return False
        key = item.get("security_code") or item.get("security_id") or item.get("ticker")
        prior = self.alert_state.get(key)
        if prior is None:
            return True
        elapsed = now - prior["time"]
        score_upgrade = float(item.get("score") or 0) >= float(prior.get("score") or 0) + 8
        crossed_call = (
            float(prior.get("score") or 0) < self.config.call_alert_score
            <= float(item.get("score") or 0)
        )
        return crossed_call or (elapsed >= timedelta(minutes=self.config.alert_cooldown_minutes) and score_upgrade)

    def _send_alert_if_needed(self, item: Dict[str, Any], now: datetime) -> None:
        if not self._should_alert(item, now):
            return
        logs = send_candidate_alerts(item, self.config, now=now)
        key = item.get("security_code") or item.get("security_id") or item.get("ticker")
        self.alert_state[key] = {"time": now, "score": item.get("score")}
        log_row = {
            "alert_time": now.isoformat(),
            "security_code": item.get("security_code"),
            "security_id": item.get("security_id"),
            "name": item.get("name"),
            "score": item.get("score"),
            "status": item.get("status"),
            "channels": "; ".join(logs),
        }
        self.alert_log.append(log_row)
        print(f"[Afternoon ALERT] {item.get('security_id')} score={item.get('score')} | {' | '.join(logs)}")

    def scan_once(self, now: Optional[datetime] = None) -> Tuple[List[Dict[str, Any]], List[Dict[str, Any]]]:
        tz = ZoneInfo(self.config.timezone)
        now = now.astimezone(tz) if now and now.tzinfo else (now.replace(tzinfo=tz) if now else datetime.now(tz))
        print(f"\n=== Afternoon fresh-move scan @ {now.strftime('%Y-%m-%d %H:%M:%S %Z')} ===")
        broad = self.broad_scan(now)
        qualified, rejected = self.deep_scan(broad, now)
        for rank, item in enumerate(qualified, start=1):
            item["rank"] = rank
            item["scan_time"] = now.isoformat()
            self._send_alert_if_needed(item, now)
        for item in rejected[:25]:
            item["scan_time"] = now.isoformat()
        self.all_snapshots.extend(_flatten_candidate(x) for x in qualified)
        self.write_outputs(now, qualified, rejected)
        print(f"[Afternoon] final: {len(qualified)} reportable | {len(rejected)} rejected deep candidates")
        for item in qualified[:10]:
            print(
                f"  {item.get('security_id'):12} score={item.get('score'):5.1f} "
                f"move={item.get('window_move_pct'):5.2f}% rvol={item.get('same_window_rvol')} "
                f"{item.get('status')} | {item.get('catalyst_strength')}"
            )
        return qualified, rejected

    def write_outputs(self, now: datetime, qualified: List[Dict[str, Any]], rejected: List[Dict[str, Any]]) -> None:
        day = now.strftime("%Y-%m-%d")
        xlsx_path = self.output_dir / f"afternoon_fresh_moves_{day}.xlsx"
        json_path = self.output_dir / f"afternoon_fresh_moves_{day}.json"

        latest_df = pd.DataFrame([_flatten_candidate(x) for x in qualified], columns=LATEST_COLUMNS)
        snapshots_df = pd.DataFrame(self.all_snapshots, columns=LATEST_COLUMNS)
        rejected_rows = []
        for item in rejected[:50]:
            row = _flatten_candidate(item)
            row["status"] = "REJECT"
            rejected_rows.append(row)
        rejected_df = pd.DataFrame(rejected_rows, columns=LATEST_COLUMNS)
        alerts_df = pd.DataFrame(self.alert_log)

        with pd.ExcelWriter(xlsx_path, engine="openpyxl", mode="w") as writer:
            latest_df.to_excel(writer, sheet_name="Latest", index=False)
            snapshots_df.to_excel(writer, sheet_name="All_Snapshots", index=False)
            rejected_df.to_excel(writer, sheet_name="Rejected_Top", index=False)
            alerts_df.to_excel(writer, sheet_name="Alert_Log", index=False)
        self._format_workbook(xlsx_path)

        payload = {
            "generated_at": now.isoformat(),
            "universe_size": len(self.universe),
            "qualified": [{k: v for k, v in x.items() if not k.startswith("_")} for x in qualified],
            "alert_log": self.alert_log,
        }
        with open(json_path, "w", encoding="utf-8") as handle:
            json.dump(payload, handle, ensure_ascii=False, indent=2, default=str)

    @staticmethod
    def _format_workbook(path: Path) -> None:
        wb = load_workbook(path)
        header_fill = PatternFill("solid", fgColor="1F3864")
        header_font = Font(color="FFFFFF", bold=True)
        for ws in wb.worksheets:
            ws.freeze_panes = "A2"
            ws.sheet_view.showGridLines = False
            for cell in ws[1]:
                cell.fill = header_fill
                cell.font = header_font
                cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
            for col_idx in range(1, ws.max_column + 1):
                header = str(ws.cell(1, col_idx).value or "")
                width = 18
                if header in ("name", "catalyst_reason", "catalyst_evidence", "reasons", "warnings", "hard_rejects"):
                    width = 42
                ws.column_dimensions[get_column_letter(col_idx)].width = width
        wb.save(path)


def wait_until_window_and_watch(scanner: AfternoonScanner) -> None:
    cfg = scanner.config
    tz = ZoneInfo(cfg.timezone)
    start_h, start_m = (int(x) for x in cfg.window_start.split(":"))
    end_h, end_m = (int(x) for x in cfg.window_end.split(":"))

    now = datetime.now(tz)
    if now.weekday() >= 5:
        print("[Afternoon] Weekend: scanner exits without running.")
        return
    start_dt = now.replace(hour=start_h, minute=start_m, second=0, microsecond=0)
    end_dt = now.replace(hour=end_h, minute=end_m, second=0, microsecond=0)
    if now < start_dt:
        seconds = (start_dt - now).total_seconds()
        print(f"[Afternoon] Waiting until {cfg.window_start} IST ({seconds/60:.1f} minutes)")
        time_module.sleep(max(0, seconds))
    elif now > end_dt:
        print(f"[Afternoon] Current time is after {cfg.window_end} IST; watch mode exits.")
        return

    next_tick = max(datetime.now(tz), start_dt)
    while True:
        now = datetime.now(tz)
        if now > end_dt:
            break
        if now < next_tick:
            time_module.sleep((next_tick - now).total_seconds())
            now = datetime.now(tz)
        try:
            scanner.scan_once(now)
        except Exception as exc:
            # A transient Yahoo/BSE failure should not kill the full 70-minute
            # watch. The next cycle gets another chance.
            print(f"[Afternoon] cycle failed but watch continues: {type(exc).__name__}: {exc}")
        next_tick += timedelta(minutes=cfg.scan_interval_minutes)
        if next_tick <= datetime.now(tz):
            next_tick = datetime.now(tz) + timedelta(seconds=20)

    print(f"[Afternoon] Window complete at {datetime.now(tz).strftime('%H:%M:%S %Z')}")


def parse_args(argv: Optional[Sequence[str]] = None) -> argparse.Namespace:
    parser = argparse.ArgumentParser(description="BSE 2PM-3:10PM fresh big-move scanner")
    mode = parser.add_mutually_exclusive_group(required=False)
    mode.add_argument("--watch", action="store_true", help="Run repeatedly from 2:00 PM through 3:10 PM IST")
    mode.add_argument("--once", action="store_true", help="Run one snapshot (also works after market close for today's 2-3:10 window)")
    parser.add_argument("--group-a", default=DEFAULT_GROUP_A)
    parser.add_argument("--group-b", default=DEFAULT_GROUP_B)
    parser.add_argument("--smallcap", default=DEFAULT_SMALLCAP)
    parser.add_argument("--output-dir", default=DEFAULT_OUTPUT_DIR)
    return parser.parse_args(argv)


def main(argv: Optional[Sequence[str]] = None) -> int:
    args = parse_args(argv)
    universe = load_afternoon_universe(args.group_a, args.group_b, args.smallcap)
    a_count = sum(1 for x in universe if x.get("group") == "A")
    b_count = sum(1 for x in universe if x.get("group") == "B")
    s_count = sum(1 for x in universe if x.get("is_smallcap"))
    print(f"[Afternoon] universe: {len(universe)} unique | Group A={a_count} | Group B={b_count} | Smallcap-tagged={s_count}")
    scanner = AfternoonScanner(universe, args.output_dir, DEFAULT_AFTERNOON_CONFIG)
    if args.watch:
        wait_until_window_and_watch(scanner)
    else:
        scanner.scan_once()
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
