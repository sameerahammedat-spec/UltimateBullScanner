"""Excel report writer for the Early Swing Rally Scanner."""
from __future__ import annotations

from datetime import date
from typing import Iterable, List, Sequence

from openpyxl import load_workbook, Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

from early_swing_models import EarlySwingCandidate, MarketRegime


FONT_NAME = "Arial"
HEADER_FILL = PatternFill("solid", fgColor="1F3864")
SECTION_FILL = PatternFill("solid", fgColor="548235")
GOOD_FILL = PatternFill("solid", fgColor="C6E0B4")
WATCH_FILL = PatternFill("solid", fgColor="FFE699")
AVOID_FILL = PatternFill("solid", fgColor="F8CBAD")
ALT_FILL = PatternFill("solid", fgColor="F2F2F2")
THIN = Side(style="thin", color="BFBFBF")
BORDER = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)

HEADERS = [
    "Rank", "Stock Name", "Symbol / Security Code", "BSE Group", "Sector", "Sources",
    "Current Price", "Setup Date", "Rally Age", "Setup Type", "Rules Score", "Confidence",
    "Suggested Status", "Entry From", "Entry To", "Stop Loss", "Target 1", "Target 2",
    "Reward:Risk", "Stop Distance %", "Position Size", "Nearest Resistance",
    "Distance to Resistance %", "1D Return %", "2D Return %", "5D Return %", "10D Return %",
    "Breakout Volume Ratio", "Current Volume Ratio", "RSI", "ADX", "ATR %", "MACD Histogram",
    "RS 2D", "RS 5D", "RS 10D", "Market Regime", "Extension Risk", "False-Breakout Risk",
    "Delivery", "Reasons Selected", "Warning Flags", "Component Scores", "Penalties",
]


def _num(value, digits=2):
    return "N/A" if value is None else round(float(value), digits)


def _row(candidate: EarlySwingCandidate, rank: int) -> list:
    return [
        rank, candidate.company_name, candidate.symbol, candidate.bse_group, candidate.sector,
        ", ".join(candidate.source_universes), _num(candidate.current_price), candidate.setup_date,
        candidate.rally_age if candidate.rally_age is not None else "N/A", candidate.setup_type,
        _num(candidate.rules_score, 1), candidate.confidence, candidate.suggested_status,
        _num(candidate.entry_low), _num(candidate.entry_high), _num(candidate.stop_loss),
        _num(candidate.target1), _num(candidate.target2), _num(candidate.reward_risk),
        _num(candidate.stop_distance_pct), candidate.position_size if candidate.position_size is not None else "N/A",
        _num(candidate.nearest_resistance), _num(candidate.distance_to_resistance_pct),
        _num(candidate.one_day_return_pct), _num(candidate.two_day_return_pct),
        _num(candidate.five_day_return_pct), _num(candidate.ten_day_return_pct),
        _num(candidate.breakout_volume_ratio), _num(candidate.current_volume_ratio),
        _num(candidate.rsi, 1), _num(candidate.adx, 1), _num(candidate.atr_pct),
        _num(candidate.macd_histogram, 3), _num(candidate.relative_strength_2d),
        _num(candidate.relative_strength_5d), _num(candidate.relative_strength_10d),
        candidate.market_regime, _num(candidate.extension_risk, 1),
        _num(candidate.false_breakout_risk, 1), candidate.delivery_status,
        "; ".join(candidate.reasons), "; ".join(candidate.warning_flags),
        "; ".join(f"{k}={v}" for k, v in candidate.component_scores.items()),
        "; ".join(candidate.penalties),
    ]


def _write_section(ws, start_row: int, title: str, rows: Sequence[EarlySwingCandidate]) -> int:
    row = start_row
    ws.cell(row=row, column=1, value=f"{title} — {len(rows)}").font = Font(name=FONT_NAME, size=11, bold=True, color="FFFFFF")
    for col in range(1, len(HEADERS) + 1):
        ws.cell(row=row, column=col).fill = SECTION_FILL
    ws.merge_cells(start_row=row, start_column=1, end_row=row, end_column=len(HEADERS))
    row += 1
    for col, header in enumerate(HEADERS, 1):
        cell = ws.cell(row=row, column=col, value=header)
        cell.font = Font(name=FONT_NAME, size=9, bold=True, color="FFFFFF")
        cell.fill = HEADER_FILL
        cell.border = BORDER
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    ws.row_dimensions[row].height = 45
    row += 1
    if not rows:
        ws.cell(row=row, column=1, value="No stocks in this category").font = Font(name=FONT_NAME, italic=True)
        return row + 2
    for idx, candidate in enumerate(rows, 1):
        values = _row(candidate, idx)
        for col, value in enumerate(values, 1):
            cell = ws.cell(row=row, column=col, value=value)
            cell.font = Font(name=FONT_NAME, size=9)
            cell.border = BORDER
            cell.alignment = Alignment(horizontal="left" if col in (2, 5, 6, 41, 42, 43, 44) else "center", wrap_text=True)
            if idx % 2 == 0:
                cell.fill = ALT_FILL
        status_cell = ws.cell(row=row, column=13)
        if candidate.suggested_status == "Enter":
            status_cell.fill = GOOD_FILL
        elif candidate.suggested_status in {"Watch", "Wait for Pullback"}:
            status_cell.fill = WATCH_FILL
        else:
            status_cell.fill = AVOID_FILL
        row += 1
    return row + 1


def write_early_swing_report(xlsx_path: str, candidates: Sequence[EarlySwingCandidate],
                             regime: MarketRegime, report_date: date) -> str:
    try:
        workbook = load_workbook(xlsx_path)
    except FileNotFoundError:
        workbook = Workbook()
        if workbook.active:
            workbook.remove(workbook.active)
    sheet_name = f"EarlySwing_{report_date.strftime('%d%b%y')}"
    if sheet_name in workbook.sheetnames:
        del workbook[sheet_name]
    ws = workbook.create_sheet(sheet_name, 0)
    ws.sheet_view.showGridLines = False
    ws.cell(row=1, column=1, value=f"Early Swing Rally Scanner — {report_date.strftime('%d %b %Y')}").font = \
        Font(name=FONT_NAME, size=14, bold=True, color="1F3864")
    ws.merge_cells(start_row=1, start_column=1, end_row=1, end_column=len(HEADERS))

    actionable = sorted([c for c in candidates if not c.rejected and c.suggested_status == "Enter"], key=lambda c: c.rules_score, reverse=True)
    pullback = sorted([c for c in candidates if not c.rejected and c.suggested_status == "Wait for Pullback"], key=lambda c: c.rules_score, reverse=True)
    watch = sorted([c for c in candidates if not c.rejected and c.suggested_status == "Watch"], key=lambda c: c.rules_score, reverse=True)
    extended = [c for c in candidates if c.rejection_category == "EXTENDED"]
    false_breakout = [c for c in candidates if c.rejection_category == "FALSE_BREAKOUT"]
    poor_rr = [c for c in candidates if c.rejection_category == "POOR_RISK_REWARD"]
    insufficient = [c for c in candidates if c.rejection_category == "INSUFFICIENT_DATA"]
    other = [c for c in candidates if c.rejected and c.rejection_category not in {
        "EXTENDED", "FALSE_BREAKOUT", "POOR_RISK_REWARD", "INSUFFICIENT_DATA"
    }]

    row = 3
    row = _write_section(ws, row, "ACTIONABLE EARLY SWING SETUPS", actionable)
    row = _write_section(ws, row, "STRONG SETUPS — WAIT FOR PULLBACK", pullback)
    row = _write_section(ws, row, "EARLY SETUPS — WATCH / NEED CONFIRMATION", watch)
    row = _write_section(ws, row, "REJECTED — EXTENDED RALLIES", extended)
    row = _write_section(ws, row, "REJECTED — FALSE BREAKOUT RISK", false_breakout)
    row = _write_section(ws, row, "REJECTED — POOR RISK/REWARD", poor_rr)
    row = _write_section(ws, row, "INSUFFICIENT / STALE DATA", insufficient)
    row = _write_section(ws, row, "OTHER REJECTIONS", other)

    ws.cell(row=row, column=1, value="MARKET DASHBOARD").font = Font(name=FONT_NAME, bold=True, color="FFFFFF")
    ws.cell(row=row, column=1).fill = SECTION_FILL
    ws.merge_cells(start_row=row, start_column=1, end_row=row, end_column=4)
    row += 1
    dashboard = [
        ("Market regime", regime.classification),
        ("Benchmark above EMA20", "Y" if regime.above_ema20 else "N"),
        ("Benchmark above SMA50", "Y" if regime.above_sma50 else "N"),
        ("Benchmark 5D return %", round(regime.five_day_return_pct, 2)),
        ("Annualised benchmark volatility %", round(regime.volatility_pct, 2)),
        ("Total analysed", len(candidates)),
        ("Actionable", len(actionable)),
    ]
    for metric, value in dashboard:
        ws.cell(row=row, column=1, value=metric).font = Font(name=FONT_NAME, bold=True)
        ws.cell(row=row, column=2, value=value)
        row += 1

    widths = {1: 8, 2: 32, 3: 18, 4: 12, 5: 20, 6: 24}
    for col in range(1, len(HEADERS) + 1):
        ws.column_dimensions[get_column_letter(col)].width = widths.get(col, 14)
    for col in (41, 42, 43, 44):
        ws.column_dimensions[get_column_letter(col)].width = 42
    ws.freeze_panes = "G5"
    workbook.save(xlsx_path)
    return sheet_name
