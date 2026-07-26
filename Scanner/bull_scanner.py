"""
Python-native Bull Scanner — Nifty 500 + Nifty Smallcap 250 universe.

Creates a NEW dated sheet every time you run it (e.g. "Scan_25Jul26"), so you
build a daily historical archive automatically. Each sheet has two sections:
  1. FULL SCAN — every stock that passed the liquidity filter, scored & ranked
  2. EARLY MOMENTUM WATCHLIST — stocks with RSI 40-50 (just turning up,
     not yet extended) filtered from the same scan, sorted by relative strength

Setup (once):
    pip install yfinance openpyxl pandas

Universe files (free, official NSE lists, no login):
    Nifty 500:          https://archives.nseindia.com/content/indices/ind_nifty500list.csv
    Nifty Smallcap 250: https://archives.nseindia.com/content/indices/ind_niftysmallcap250list.csv
    Save both in the same folder as this script.

Usage:
    python bull_scanner.py trading_toolkit.xlsx nifty500.csv niftysmallcap250.csv
"""

import sys
import time
import pickle
from datetime import date
import pandas as pd
import yfinance as yf
from openpyxl import load_workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.formatting.rule import CellIsRule
from openpyxl.utils import get_column_letter

FONT = "Arial"
BLACK = Font(name=FONT, size=10)
BLACK_B = Font(name=FONT, size=10, bold=True)
HDR_FONT = Font(name=FONT, color="FFFFFF", size=10, bold=True)
HDR_FILL = PatternFill("solid", fgColor="1F3864")
SEC_FILL = PatternFill("solid", fgColor="548235")
SEC_FONT = Font(name=FONT, color="FFFFFF", size=11, bold=True)
ALT_FILL = PatternFill("solid", fgColor="F2F2F2")
GOOD_FILL = PatternFill("solid", fgColor="C6E0B4")
WATCH_FILL = PatternFill("solid", fgColor="FFE699")
AVOID_FILL = PatternFill("solid", fgColor="F8CBAD")
thin = Side(style="thin", color="BFBFBF")
BORDER = Border(left=thin, right=thin, top=thin, bottom=thin)

MIN_VOLUME = 50000
NIFTY_TICKER = "^NSEI"

HEADERS = ["Symbol", "Name", "Sector", "Close", "% Chg", "Volume", "Mkt Cap (Cr)",
           "Liquidity OK?", "Above 50DMA?", "Above 200DMA?", "RSI(14)", "RSI Healthy?",
           "ATR(14)", "Vol Spike?", "% Down from 52wH", "Near 52wH?",
           "RS vs Nifty 1M %", "RS vs Nifty 3M %", "RS Positive?",
           "P/E", "ROE %", "Debt/Equity", "QoQ Profit Gr %",
           "Score (0-9)", "Verdict"]


def rsi(series, period=14):
    delta = series.diff()
    gain = delta.clip(lower=0)
    loss = -delta.clip(upper=0)
    avg_gain = gain.rolling(period).mean()
    avg_loss = loss.rolling(period).mean()
    rs = avg_gain / avg_loss
    return 100 - (100 / (1 + rs))


def atr(df, period=14):
    high, low, close = df["High"], df["Low"], df["Close"]
    prev_close = close.shift(1)
    tr = pd.concat([(high - low), (high - prev_close).abs(), (low - prev_close).abs()], axis=1).max(axis=1)
    return tr.rolling(period).mean()


def to_scalar(x):
    if isinstance(x, pd.Series):
        x = x.iloc[-1] if len(x) else None
    if x is None or pd.isna(x):
        return None
    return float(x)


def get_nifty_returns():
    tk = yf.Ticker(NIFTY_TICKER)
    df = tk.history(period="1y", interval="1d", auto_adjust=True)
    close = df["Close"]
    ret_1m = (close.iloc[-1] / close.iloc[-22] - 1) * 100 if len(close) > 22 else None
    ret_3m = (close.iloc[-1] / close.iloc[-63] - 1) * 100 if len(close) > 63 else None
    return ret_1m, ret_3m


def analyze_symbol(nse_symbol, nifty_1m, nifty_3m):
    ticker_str = nse_symbol.strip() + ".NS"
    try:
        tk = yf.Ticker(ticker_str)
        df = tk.history(period="1y", interval="1d", auto_adjust=True)
        if df.empty or len(df) < 60:
            return None
    except Exception:
        return None

    if isinstance(df.columns, pd.MultiIndex):
        df.columns = df.columns.get_level_values(0)

    close, volume = df["Close"], df["Volume"]
    last_close = to_scalar(close.iloc[-1])
    prev_close = to_scalar(close.iloc[-2]) if len(close) > 1 else None
    pct_change = ((last_close - prev_close) / prev_close * 100) if (last_close and prev_close) else None

    dma50 = to_scalar(close.rolling(50).mean().iloc[-1])
    dma200 = to_scalar(close.rolling(200).mean().iloc[-1]) if len(close) >= 200 else None
    rsi14 = to_scalar(rsi(close, 14).iloc[-1])
    atr14 = to_scalar(atr(df, 14).iloc[-1])

    avg_vol_20 = to_scalar(volume.rolling(20).mean().iloc[-1])
    last_vol = to_scalar(volume.iloc[-1])
    vol_spike = (last_vol is not None and avg_vol_20 and last_vol > 1.5 * avg_vol_20)

    high_52w = to_scalar(close.rolling(252).max().iloc[-1]) if len(close) >= 60 else to_scalar(close.max())
    down_from_high = ((high_52w - last_close) / high_52w * 100) if (high_52w and last_close) else None

    stock_1m = (last_close / to_scalar(close.iloc[-22]) - 1) * 100 if len(close) > 22 else None
    stock_3m = (last_close / to_scalar(close.iloc[-63]) - 1) * 100 if len(close) > 63 else None
    rs_1m = (stock_1m - nifty_1m) if (stock_1m is not None and nifty_1m is not None) else None
    rs_3m = (stock_3m - nifty_3m) if (stock_3m is not None and nifty_3m is not None) else None

    roe = pe = debt_equity = qoq_growth = mkt_cap_cr = None
    try:
        info = tk.info
        if info.get("returnOnEquity") is not None:
            roe = info["returnOnEquity"] * 100
        pe = info.get("trailingPE")
        debt_equity = info.get("debtToEquity")
        if info.get("marketCap") is not None:
            mkt_cap_cr = info["marketCap"] / 1e7
    except Exception:
        pass
    try:
        qf = tk.quarterly_financials
        if qf is not None and "Net Income" in qf.index and qf.shape[1] >= 2:
            latest_q, prev_q = qf.loc["Net Income"].iloc[0], qf.loc["Net Income"].iloc[1]
            if prev_q and prev_q != 0:
                qoq_growth = (latest_q - prev_q) / abs(prev_q) * 100
    except Exception:
        pass

    return dict(
        close=last_close, pct_change=pct_change, volume=last_vol, mkt_cap_cr=mkt_cap_cr,
        liquidity_ok=(last_vol is not None and last_vol >= MIN_VOLUME),
        above_50dma=(dma50 is not None and last_close > dma50),
        above_200dma=(dma200 is not None and last_close > dma200) if dma200 is not None else None,
        rsi=rsi14, rsi_healthy=(rsi14 is not None and 50 <= rsi14 <= 70), atr=atr14,
        vol_spike=vol_spike, down_from_high=down_from_high,
        near_high=(down_from_high is not None and down_from_high <= 10),
        rs_1m=rs_1m, rs_3m=rs_3m, rs_positive=(rs_3m is not None and rs_3m > 0),
        pe=pe, roe=roe, roe_ok=(roe is not None and roe > 15), debt_equity=debt_equity,
        qoq_growth=qoq_growth, qoq_ok=(qoq_growth is not None and qoq_growth > 0),
    )


def score_row(d):
    checks = [d["liquidity_ok"], d["above_50dma"], bool(d["above_200dma"]),
              d["rsi_healthy"], d["vol_spike"], d["near_high"],
              d["rs_positive"], d["roe_ok"], d["qoq_ok"]]
    return sum(1 for c in checks if c)


def verdict(score):
    if score >= 7:
        return "STRONG SETUP"
    elif score >= 5:
        return "WATCHLIST"
    return "AVOID"


def find_col(df, names):
    for c in df.columns:
        if c.strip().lower() in names:
            return c
    return None


def load_universe(paths):
    frames = []
    for p in paths:
        df = pd.read_csv(p)
        sym_col = find_col(df, ("symbol", "nse code", "nsecode", "ticker"))
        if not sym_col:
            print(f"  [!] Skipping {p}: no Symbol column found ({list(df.columns)})")
            continue
        name_col = find_col(df, ("company name", "name", "stock name"))
        sector_col = find_col(df, ("industry", "sector"))
        out = pd.DataFrame({
            "symbol": df[sym_col].astype(str).str.strip(),
            "name": df[name_col].astype(str).str.strip() if name_col else df[sym_col],
            "sector": df[sector_col].astype(str).str.strip() if sector_col else "",
        })
        frames.append(out)
    merged = pd.concat(frames, ignore_index=True).drop_duplicates(subset="symbol")
    return merged


def write_table(ws, start_row, rows, title=None):
    r = start_row
    if title:
        ws.cell(row=r, column=1, value=title).font = SEC_FONT
        for c in range(1, len(HEADERS) + 1):
            ws.cell(row=r, column=c).fill = SEC_FILL
        ws.merge_cells(start_row=r, start_column=1, end_row=r, end_column=len(HEADERS))
        r += 1

    for i, h in enumerate(HEADERS, start=1):
        c = ws.cell(row=r, column=i, value=h)
        c.font = HDR_FONT
        c.fill = HDR_FILL
        c.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        c.border = BORDER
    ws.row_dimensions[r].height = 32
    header_row = r
    r += 1

    def yn(v):
        return "N/A" if v is None else ("Y" if v else "N")

    def num(v, nd=1):
        return round(v, nd) if v is not None else "N/A"

    for idx, d in enumerate(rows):
        vals = [d["symbol"], d["name"], d.get("sector", ""), num(d["close"], 2), num(d["pct_change"], 2),
                d["volume"], num(d["mkt_cap_cr"], 0), yn(d["liquidity_ok"]), yn(d["above_50dma"]),
                yn(d["above_200dma"]), num(d["rsi"], 1), yn(d["rsi_healthy"]), num(d["atr"], 2),
                yn(d["vol_spike"]), num(d["down_from_high"], 1), yn(d["near_high"]),
                num(d["rs_1m"], 1), num(d["rs_3m"], 1), yn(d["rs_positive"]),
                num(d["pe"], 1), num(d["roe"], 1), num(d["debt_equity"], 2), num(d["qoq_growth"], 1),
                d["score"], d["verdict"]]
        for c, v in enumerate(vals, start=1):
            cell = ws.cell(row=r, column=c, value=v)
            cell.font = BLACK_B if c == len(HEADERS) else BLACK
            cell.border = BORDER
            cell.alignment = Alignment(horizontal="center")
            if idx % 2 == 1:
                cell.fill = ALT_FILL
        r += 1

    last_row = r - 1
    verdict_col = get_column_letter(len(HEADERS))
    if last_row >= header_row + 1:
        ws.conditional_formatting.add(f"{verdict_col}{header_row+1}:{verdict_col}{last_row}",
            CellIsRule(operator="equal", formula=['"STRONG SETUP"'], fill=GOOD_FILL))
        ws.conditional_formatting.add(f"{verdict_col}{header_row+1}:{verdict_col}{last_row}",
            CellIsRule(operator="equal", formula=['"WATCHLIST"'], fill=WATCH_FILL))
        ws.conditional_formatting.add(f"{verdict_col}{header_row+1}:{verdict_col}{last_row}",
            CellIsRule(operator="equal", formula=['"AVOID"'], fill=AVOID_FILL))

    return r + 1  # next free row


def main(xlsx_path, universe_paths):
    universe = load_universe(universe_paths)
    symbols = universe["symbol"].tolist()
    meta = universe.set_index("symbol").to_dict(orient="index")

    print(f"Universe loaded: {len(symbols)} unique symbols across {len(universe_paths)} list(s)")
    print("Fetching Nifty 50 benchmark returns...")
    nifty_1m, nifty_3m = get_nifty_returns()
    print(f"Nifty 1M: {nifty_1m:.2f}%  |  Nifty 3M: {nifty_3m:.2f}%\n")

    results = []
    for i, sym in enumerate(symbols, 1):
        print(f"[{i}/{len(symbols)}] {sym}...", end=" ", flush=True)
        try:
            d = analyze_symbol(sym, nifty_1m, nifty_3m)
        except Exception as e:
            print(f"SKIPPED ({e})")
            continue
        if d is None:
            print("SKIPPED (no data)")
            continue
        if not d["liquidity_ok"]:
            print("SKIPPED (below liquidity cutoff)")
            continue
        d["symbol"] = sym
        d["name"] = meta.get(sym, {}).get("name", sym)
        d["sector"] = meta.get(sym, {}).get("sector", "")
        d["score"] = score_row(d)
        d["verdict"] = verdict(d["score"])
        results.append(d)
        print(f"Score={d['score']} -> {d['verdict']}")
        time.sleep(0.25)

    full_scan = sorted(results, key=lambda x: x["score"], reverse=True)
    early_momentum = sorted(
        [d for d in results if d["rsi"] is not None and 40 <= d["rsi"] <= 50 and d["liquidity_ok"]],
        key=lambda x: (x["rs_3m"] if x["rs_3m"] is not None else -999), reverse=True
    )

    today = date.today()
    sheet_name = f"Scan_{today.strftime('%d%b%y')}"  # e.g. Scan_25Jul26

    backup_path = "bull_scan_backup.pkl"
    with open(backup_path, "wb") as f:
        pickle.dump({"date": today, "sheet_name": sheet_name,
                     "full_scan": full_scan, "early_momentum": early_momentum}, f)
    print(f"Backup saved to {backup_path} (safe even if the Excel write below fails)\n")

    try:
        write_to_workbook(xlsx_path, sheet_name, today, full_scan, early_momentum)
    except PermissionError:
        print(f"\n[!] Could not save {xlsx_path} — it's likely still open in Excel/LibreOffice.")
        print("    1. Close the Excel file completely.")
        print("    2. Run:  python write_backup_to_excel.py " + xlsx_path)
        print("    (This will NOT re-fetch any data — it writes from the backup instantly.)")
        sys.exit(1)

    print(f"\nDone. Sheet '{sheet_name}' created with {len(full_scan)} full-scan stocks "
          f"and {len(early_momentum)} early-momentum candidates.")
    print("P/E, ROE, Debt/Equity, QoQ Growth are best-effort from Yahoo Finance — cross-check on Screener.in before relying on them.")


def write_to_workbook(xlsx_path, sheet_name, today, full_scan, early_momentum):
    wb = load_workbook(xlsx_path)
    if sheet_name in wb.sheetnames:
        del wb[sheet_name]
    ws = wb.create_sheet(sheet_name, 1)
    ws.sheet_view.showGridLines = False

    ws.cell(row=1, column=1, value=f"Bull Scan — {today.strftime('%d %b %Y')}").font = Font(name=FONT, size=14, bold=True, color="1F3864")
    ws.merge_cells(start_row=1, start_column=1, end_row=1, end_column=len(HEADERS))

    next_row = write_table(ws, 3, full_scan, title=f"FULL SCAN — {len(full_scan)} stocks passed liquidity filter, ranked by score")
    write_table(ws, next_row, early_momentum,
                title=f"EARLY MOMENTUM WATCHLIST (RSI 40-50) — {len(early_momentum)} stocks turning up, not yet extended")

    for i in range(1, len(HEADERS) + 1):
        ws.column_dimensions[get_column_letter(i)].width = 13
    ws.column_dimensions["B"].width = 28
    ws.column_dimensions["C"].width = 18
    ws.freeze_panes = "D4"

    wb.save(xlsx_path)


if __name__ == "__main__":
    if len(sys.argv) < 3:
        print("Usage: python bull_scanner.py trading_toolkit.xlsx nifty500.csv [niftysmallcap250.csv ...]")
        sys.exit(1)
    main(sys.argv[1], sys.argv[2:])
